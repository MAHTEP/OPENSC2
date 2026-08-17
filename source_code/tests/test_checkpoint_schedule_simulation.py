from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from simulation import Simulation
from utility_functions.checkpoint_schedule import (
    CheckpointBoundary,
    CheckpointSchedule,
)


class _StopAfterPlannerCall(Exception):
    pass


class _TransientFrame:
    def __init__(self, values):
        self.values = dict(values)

    def __getitem__(self, key):
        if key != "Value":
            raise KeyError(key)
        return self

    def to_dict(self):
        return dict(self.values)


class CheckpointScheduleSimulationTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)
        self.starter_path = self.root / "transitory_input.xlsx"
        workbook = Workbook()
        transient = workbook.active
        transient.title = "TRANSIENT"
        transient["A1"] = "OPENSC2 transient input"
        transient.append(("Variable name", "Value"))
        workbook.save(self.starter_path)
        workbook.close()
        self.transient_input = {
            "SIMULATION": "schedule-test",
            "MAGNET": "conductor_definition.xlsx",
            "ENVIRONMENT": "environment_input.xlsx",
            "TEND": 0.65,
            "IADAPTIME": 0,
            "STPMIN": 0.01,
            "MLT_INCREASE": 1.2,
            "MLT_DECREASE": 0.5,
            "USER_CHECKPOINTS": True,
        }

    def test_initialization_loads_and_stores_checkpoint_schedule(self):
        expected_schedule = SimpleNamespace(
            user_enabled=True,
            boundaries=(),
        )
        call_order = []

        def load_schedule(*args, **kwargs):
            call_order.append("checkpoint_schedule")
            return expected_schedule

        def make_environment(*args, **kwargs):
            call_order.append("environment")
            return SimpleNamespace()

        with (
            patch.object(
                Simulation,
                "CWD",
                str(self.root / "source_code"),
            ),
            patch(
                "simulation.pd.read_excel",
                return_value=_TransientFrame(self.transient_input),
            ),
            patch("simulation.check_flag_value"),
            patch(
                "simulation.load_checkpoint_schedule",
                side_effect=load_schedule,
            ) as loader,
            patch(
                "simulation.Environment",
                side_effect=make_environment,
            ),
        ):
            simulation = Simulation(str(self.root))

        self.assertIs(
            simulation.checkpoint_schedule,
            expected_schedule,
        )
        loader.assert_called_once_with(
            str(self.starter_path),
            simulation.transient_input,
            epsilon=simulation.epsilon,
        )
        self.assertEqual(
            call_order,
            ["checkpoint_schedule", "environment"],
        )

    def test_invalid_schedule_stops_before_environment_initialization(self):
        with (
            patch.object(
                Simulation,
                "CWD",
                str(self.root / "source_code"),
            ),
            patch(
                "simulation.pd.read_excel",
                return_value=_TransientFrame(self.transient_input),
            ),
            patch("simulation.check_flag_value"),
            patch(
                "simulation.load_checkpoint_schedule",
                side_effect=ValueError("invalid checkpoint schedule"),
            ),
            patch("simulation.Environment") as environment,
        ):
            with self.assertRaisesRegex(
                ValueError,
                "invalid checkpoint schedule",
            ):
                Simulation(str(self.root))

        environment.assert_not_called()

    def test_initialization_preserves_boolean_with_iadaptime_one(self):
        workbook = Workbook()
        transient = workbook.active
        transient.title = "TRANSIENT"
        transient["A1"] = "OPENSC2 transient input"
        transient.append(("Variable name", "Value"))
        for name, value in (
            ("SIMULATION", "boolean-type-test"),
            ("MAGNET", "conductor_definition.xlsx"),
            ("ENVIRONMENT", "environment_input.xlsx"),
            ("TEND", 0.02),
            ("IADAPTIME", 1),
            ("TIME_STEP", 0.00125),
            ("STPMIN", 0.000625),
            ("STPMAX", 0.0025),
            ("MLT_INCREASE", 1.2),
            ("MLT_DECREASE", 0.5),
            ("USER_CHECKPOINTS", True),
        ):
            transient.append((name, value))
        checkpoints = workbook.create_sheet("CHECKPOINTS")
        checkpoints["A1"] = "Time (s)"
        workbook.save(self.starter_path)
        workbook.close()

        with (
            patch.object(
                Simulation,
                "CWD",
                str(self.root / "source_code"),
            ),
            patch("simulation.check_flag_value"),
            patch(
                "simulation.Environment",
                return_value=SimpleNamespace(),
            ),
        ):
            simulation = Simulation(str(self.root))

        self.assertIs(
            simulation.transient_input["USER_CHECKPOINTS"],
            True,
        )
        self.assertIsInstance(
            simulation.transient_input["USER_CHECKPOINTS"],
            bool,
        )
        self.assertTrue(simulation.checkpoint_schedule.user_enabled)
        self.assertEqual(
            simulation.checkpoint_schedule.boundaries,
            (
                CheckpointBoundary(time=0.02, trigger="final"),
            ),
        )

    def _observe_first_planner_call(self, schedule, *, current_time=0.0):
        simulation = Simulation.__new__(Simulation)
        simulation.checkpoint_schedule = schedule
        simulation.epsilon = 1.0e-6
        simulation.simulation_time = [current_time]
        simulation.transient_input = {
            "TEND": 1.0,
            "STPMIN": 0.01,
        }
        simulation.num_step = 0
        simulation.numObj = 1
        simulation.restored_from_checkpoint = True
        simulation.environment = SimpleNamespace()

        conductor = SimpleNamespace(
            compute_radiative_heat_exhange_jk=lambda: None,
            compute_heat_exchange_jk_env=lambda environment: None,
        )
        simulation.list_of_Conductors = [conductor]
        observed = {}

        def stop_at_planner(*args, **kwargs):
            observed["args"] = args
            observed["kwargs"] = kwargs
            raise _StopAfterPlannerCall

        with patch(
            "simulation.plan_next_time_step",
            side_effect=stop_at_planner,
        ):
            with self.assertRaises(_StopAfterPlannerCall):
                simulation.conductor_solution(gui=None)

        return simulation, conductor, observed

    def test_solution_passes_next_checkpoint_boundary_to_time_planner(self):
        schedule = CheckpointSchedule(
            user_enabled=True,
            boundaries=(
                CheckpointBoundary(time=0.25, trigger="requested"),
                CheckpointBoundary(time=1.0, trigger="final"),
            ),
        )

        simulation, conductor, observed = self._observe_first_planner_call(
            schedule
        )

        self.assertIs(observed["args"][0], conductor)
        self.assertEqual(
            observed["args"][1:],
            (simulation.epsilon, 0.01),
        )
        self.assertEqual(
            observed["kwargs"],
            {"scheduled_boundary_time": 0.25},
        )

    def test_solution_passes_none_when_no_future_boundary_exists(self):
        schedules_and_times = (
            (CheckpointSchedule(user_enabled=False, boundaries=()), 0.0),
            (
                CheckpointSchedule(
                    user_enabled=True,
                    boundaries=(
                        CheckpointBoundary(
                            time=0.25,
                            trigger="requested",
                        ),
                        CheckpointBoundary(time=0.5, trigger="final"),
                    ),
                ),
                0.5,
            ),
        )

        observed_keyword_arguments = []
        for schedule, current_time in schedules_and_times:
            with self.subTest(
                user_enabled=schedule.user_enabled,
                current_time=current_time,
            ):
                _, _, observed = self._observe_first_planner_call(
                    schedule,
                    current_time=current_time,
                )
                observed_keyword_arguments.append(observed["kwargs"])

        self.assertEqual(
            observed_keyword_arguments,
            [
                {"scheduled_boundary_time": None},
                {"scheduled_boundary_time": None},
            ],
        )


if __name__ == "__main__":
    unittest.main()
