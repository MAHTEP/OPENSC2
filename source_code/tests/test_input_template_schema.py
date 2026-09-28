from pathlib import Path
import unittest

from openpyxl import load_workbook

from simulation import _read_transient_input_preserving_excel_booleans


class InputTemplateSchemaTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.template_root = (
            Path(__file__).resolve().parents[1]
            / "input_files"
            / "input_file_template"
        )

    def _row_names(self, filename, sheet_name, start_row=4):
        workbook = load_workbook(
            self.template_root / filename,
            read_only=True,
            data_only=True,
        )
        self.addCleanup(workbook.close)
        return {
            row[0]
            for row in workbook[sheet_name].iter_rows(
                min_row=start_row,
                max_col=1,
                values_only=True,
            )
            if row[0] is not None
        }

    def test_stack_uses_the_implemented_schema(self):
        row_names = self._row_names(
            "template_conductor_1_input.xlsx",
            "STACK",
        )
        self.assertTrue(
            {
                "N_tape",
                "Stack_width",
                "superconducting_material",
                "RRR",
                "RRR_Ag",
                "C0_MODE",
            }.issubset(row_names)
        )
        self.assertTrue(
            {"Tape_number", "Stack_witdh", "HTS_material"}.isdisjoint(
                row_names
            )
        )

    def test_cryosoft_inputs_are_documented_where_used(self):
        stack_rows = self._row_names(
            "template_conductor_1_input.xlsx",
            "STACK",
        )
        jacket_rows = self._row_names(
            "template_conductor_1_input.xlsx",
            "Z_JACKET",
        )
        stabilizer_rows = self._row_names(
            "template_conductor_1_input.xlsx",
            "STR_STAB",
        )
        self.assertIn("RRR_Ag", stack_rows)
        self.assertIn("RRR", jacket_rows)
        self.assertIn("RRR", stabilizer_rows)

    def test_operation_schema_exposes_only_supported_controls(self):
        filename = "template_conductor_1_operation.xlsx"
        for sheet_name in ("STACK", "STR_MIX", "STR_STAB", "Z_JACKET"):
            with self.subTest(sheet_name=sheet_name):
                row_names = self._row_names(filename, sheet_name)
                self.assertNotIn("fixAlphaBvalue", row_names)
                self.assertNotIn("QJFRACT", row_names)
                self.assertNotIn("BITR", row_names)
                self.assertNotIn("BOTR", row_names)

        for sheet_name in ("STACK", "STR_MIX"):
            with self.subTest(sheet_name=sheet_name):
                self.assertIn(
                    "EPS_INTERPOLATION",
                    self._row_names(filename, sheet_name),
                )

        channel_rows = self._row_names(filename, "CHAN")
        self.assertTrue({"TEMINI", "PREINI"}.isdisjoint(channel_rows))

    def test_obsolete_conductor_fields_are_absent(self):
        filename = "template_conductor_definition.xlsx"
        input_rows = self._row_names(filename, "CONDUCTOR_input")
        self.assertTrue(
            {
                "ISJOINT",
                "XJBEG",
                "XJBEIN",
                "XJBEOUT",
                "XJENOUT",
                "external_free_convection_correlation",
            }.isdisjoint(input_rows)
        )
        operation_rows = self._row_names(filename, "CONDUCTOR_operation")
        self.assertIn("MAXIMUM_ITERATION_NUMBER", operation_rows)

    def test_transverse_transport_multiplier_is_available(self):
        workbook = load_workbook(
            self.template_root / "template_conductor_1_coupling.xlsx",
            read_only=True,
            data_only=True,
        )
        self.addCleanup(workbook.close)
        self.assertIn("trans_transp_multiplier", workbook.sheetnames)
        worksheet = workbook["trans_transp_multiplier"]
        self.assertEqual(worksheet["D4"].value, 1)

    def test_archived_transient_controls_are_absent(self):
        transient_rows = self._row_names(
            "template_transitory_input.xlsx",
            "TRANSIENT",
            start_row=3,
        )
        self.assertTrue({"TIMEREF", "TAUREF"}.isdisjoint(transient_rows))

    def test_runtime_loader_reads_the_transient_template(self):
        transient_input = _read_transient_input_preserving_excel_booleans(
            self.template_root / "template_transitory_input.xlsx"
        )

        self.assertEqual(transient_input["SIMULATION"], "opensc2_template")
        self.assertIs(transient_input["USER_CHECKPOINTS"], False)

    def test_templates_do_not_contain_external_workbook_formulas(self):
        for path in sorted(self.template_root.glob("*.xlsx")):
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=False,
            )
            self.addCleanup(workbook.close)
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        value = cell.value
                        if isinstance(value, str) and value.startswith("="):
                            with self.subTest(
                                filename=path.name,
                                sheet=worksheet.title,
                                cell=cell.coordinate,
                            ):
                                self.assertNotIn("[", value)

    def test_cached_headers_are_available_to_opensc2(self):
        workbook = load_workbook(
            self.template_root / "template_conductor_1_input.xlsx",
            read_only=True,
            data_only=True,
        )
        self.addCleanup(workbook.close)
        expected = {
            "CHAN": (2, "CHAN_1"),
            "STACK": (0, "STACK_0"),
            "STR_MIX": (1, "STR_MIX_1"),
            "STR_STAB": (0, "STR_STAB_0"),
            "Z_JACKET": (1, "Z_JACKET_1"),
        }
        for sheet_name, (count, identifier) in expected.items():
            with self.subTest(sheet_name=sheet_name):
                worksheet = workbook[sheet_name]
                self.assertEqual(worksheet["B1"].value, count)
                self.assertEqual(worksheet["E3"].value, identifier)


if __name__ == "__main__":
    unittest.main()
