import hashlib
from pathlib import Path
import stat
import tempfile
import unittest

from openpyxl import load_workbook, Workbook

from simulation import Simulation


class CheckpointInputMetadataTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)

        self.root = Path(self.temporary_directory.name)
        self.input_directory = self.root / "inputs"
        self.metadata_directory = self.root / "metadata"
        self.input_directory.mkdir()
        self.metadata_directory.mkdir()

        self.input_names = (
            "transitory_input.xlsx",
            "environment_input.xlsx",
            "conductor_definition.xlsx",
            "conductor_input.xlsx",
            "conductor_coupling.xlsx",
            "auxiliary_profile.tsv",
        )
        self._build_input_family()

        self.simulation = Simulation.__new__(Simulation)
        self.simulation.basePath = str(self.input_directory)
        self.simulation.starter_file = "transitory_input.xlsx"
        self.simulation.starter_file_path = str(
            self.input_directory / self.simulation.starter_file
        )
        self.simulation.dict_path = {
            "Save_input": str(self.metadata_directory)
        }

    def _write_workbook(self, filename, sheets):
        workbook = Workbook()
        workbook.remove(workbook.active)

        for sheet_name, rows in sheets.items():
            worksheet = workbook.create_sheet(sheet_name)
            for row in rows:
                worksheet.append(row)

        path = self.input_directory / filename
        workbook.save(path)
        workbook.close()

    def _write_binary_file(self, filename, content):
        path = self.input_directory / filename
        path.write_bytes(content)

    def _build_input_family(self):
        self._write_workbook(
            "transitory_input.xlsx",
            {
                "TRANSIENT": (
                    ("% master transitory input",),
                    ("Variable name", "Value"),
                    ("ENVIRONMENT", "environment_input.xlsx"),
                    ("MAGNET", "conductor_definition.xlsx"),
                ),
                "CHECKPOINTS": (
                    ("Time (s)",),
                    (0.125,),
                    (0.375,),
                ),
            },
        )
        self._write_workbook(
            "environment_input.xlsx",
            {
                "ENVIRONMENT": (
                    ("% environment input",),
                    ("Variable name", "Value"),
                    ("Temperature", 4.5),
                ),
            },
        )
        self._write_workbook(
            "conductor_definition.xlsx",
            {
                "CONDUCTOR_files": (
                    ("% conductor definition",),
                    ("% second descriptive row",),
                    (
                        "Variable name",
                        "Unit",
                        "Variable type",
                        "Note/comments",
                        "CONDUCTOR_1",
                    ),
                    (
                        "STRUCTURE_INPUT",
                        "-",
                        "string",
                        "",
                        "conductor_input.xlsx",
                    ),
                    (
                        "STRUCTURE_COUPLING",
                        "-",
                        "string",
                        "",
                        "conductor_coupling.xlsx",
                    ),
                    (
                        "EXTERNAL_PROFILE",
                        "-",
                        "string",
                        "",
                        "auxiliary_profile.tsv",
                    ),
                ),
            },
        )
        self._write_workbook(
            "conductor_input.xlsx",
            {
                "CHAN": (
                    ("% conductor input",),
                    ("% second descriptive row",),
                    ("Variable name", "CHAN_1"),
                    ("HYDIAMETER", 0.01),
                ),
            },
        )
        self._write_workbook(
            "conductor_coupling.xlsx",
            {
                "contact_perimeter": (
                    ("% coupling input",),
                    ("Component", "CHAN_1"),
                    ("CHAN_1", 0.0),
                ),
            },
        )
        self._write_binary_file(
            "auxiliary_profile.tsv",
            b"z (m)\tvalue\n0.0\t1.0\n1.0\t2.0\n",
        )

    @staticmethod
    def _sha256(path):
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for block in iter(lambda: stream.read(65536), b""):
                digest.update(block)
        return digest.hexdigest()

    def test_metadata_preserves_every_input_file_byte_for_byte(self):
        source_hashes = {
            name: self._sha256(self.input_directory / name)
            for name in self.input_names
        }

        self.simulation.save_input_files()

        problems = []
        for name in self.input_names:
            destination = self.metadata_directory / f"meta_{name}"
            if not destination.is_file():
                problems.append(f"{name}: metadata copy is missing")
                continue

            destination_hash = self._sha256(destination)
            if destination_hash != source_hashes[name]:
                problems.append(
                    f"{name}: metadata copy differs from its source"
                )

            if destination.stat().st_mode & stat.S_IWUSR:
                problems.append(
                    f"{name}: metadata copy is still owner-writable"
                )

        self.assertEqual(problems, [])

        for name in self.input_names:
            self.assertEqual(
                self._sha256(self.input_directory / name),
                source_hashes[name],
                f"The source input file {name!r} was modified.",
            )

    def test_metadata_preserves_checkpoint_sheet_header_and_first_time(self):
        self.simulation.save_input_files()

        source_path = self.input_directory / "transitory_input.xlsx"
        metadata_path = (
            self.metadata_directory / "meta_transitory_input.xlsx"
        )
        source = load_workbook(
            source_path,
            read_only=True,
            data_only=False,
        )
        metadata = load_workbook(
            metadata_path,
            read_only=True,
            data_only=False,
        )
        self.addCleanup(source.close)
        self.addCleanup(metadata.close)

        self.assertEqual(metadata.sheetnames, source.sheetnames)
        self.assertEqual(
            [
                metadata["CHECKPOINTS"]["A1"].value,
                metadata["CHECKPOINTS"]["A2"].value,
                metadata["CHECKPOINTS"]["A3"].value,
            ],
            ["Time (s)", 0.125, 0.375],
        )

    def test_metadata_overwrites_existing_read_only_copy(self):
        self.simulation.save_input_files()

        source_path = self.input_directory / "auxiliary_profile.tsv"
        metadata_path = (
            self.metadata_directory / "meta_auxiliary_profile.tsv"
        )
        updated_content = b"z (m)\tvalue\n0.0\t3.0\n1.0\t4.0\n"
        source_path.write_bytes(updated_content)

        self.simulation.save_input_files()

        self.assertEqual(metadata_path.read_bytes(), updated_content)
        self.assertFalse(metadata_path.stat().st_mode & stat.S_IWUSR)


if __name__ == "__main__":
    unittest.main()
