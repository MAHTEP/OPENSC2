from pathlib import Path
import unittest

from openpyxl import load_workbook

from utility_functions.checkpoint import checkpoint_interval
from utility_functions.checkpoint_schedule import load_checkpoint_schedule


class CheckpointInputTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.template_path = (
            Path(__file__).resolve().parents[1]
            / "input_files"
            / "input_file_template"
            / "template_transitory_input.xlsx"
        )
        if not cls.template_path.is_file():
            raise FileNotFoundError(
                f"Input template not found: {cls.template_path!s}"
            )

    def setUp(self):
        self.workbook = load_workbook(
            self.template_path,
            read_only=True,
            data_only=False,
        )
        self.addCleanup(self.workbook.close)

    def _transient_rows(self):
        worksheet = self.workbook["TRANSIENT"]
        rows = {}
        for row in worksheet.iter_rows(
            min_row=3,
            max_col=5,
            values_only=True,
        ):
            variable_name = row[0]
            if variable_name is not None:
                self.assertNotIn(
                    variable_name,
                    rows,
                    f"Duplicate TRANSIENT variable {variable_name!r}.",
                )
                rows[variable_name] = row
        return rows

    def test_template_declares_explicit_checkpoint_controls(self):
        rows = self._transient_rows()
        expected_controls = {
            "CHECKPOINT_EVERY_N_STEPS",
            "USER_CHECKPOINTS",
        }
        missing_controls = expected_controls.difference(rows)
        self.assertFalse(
            missing_controls,
            "Missing checkpoint controls in TRANSIENT: "
            f"{sorted(missing_controls)!r}",
        )

        periodic_row = rows["CHECKPOINT_EVERY_N_STEPS"]
        self.assertEqual(periodic_row[1], "-")
        self.assertEqual(periodic_row[2].strip().casefold(), "integer")
        self.assertEqual(periodic_row[4], 100)
        self.assertIs(type(periodic_row[4]), int)

        user_row = rows["USER_CHECKPOINTS"]
        self.assertEqual(user_row[1], "-")
        self.assertEqual(user_row[2].strip().casefold(), "boolean")
        self.assertIs(user_row[4], False)
        user_note = user_row[3]
        self.assertIsInstance(user_note, str)
        normalized_note = user_note.casefold()
        self.assertIn("column a", normalized_note)
        self.assertIn("a2", normalized_note)
        self.assertIn("time (s)", normalized_note)

        transient_input = {
            name: row[4]
            for name, row in rows.items()
        }
        self.assertEqual(checkpoint_interval(transient_input), 100)
        schedule = load_checkpoint_schedule(
            self.template_path,
            transient_input,
        )
        self.assertFalse(schedule.user_enabled)
        self.assertEqual(schedule.boundaries, ())

    def test_template_contains_empty_checkpoints_sheet(self):
        self.assertIn("CHECKPOINTS", self.workbook.sheetnames)
        worksheet = self.workbook["CHECKPOINTS"]
        self.assertEqual(worksheet["A1"].value, "Time (s)")

        payload = [
            cell.value
            for row in worksheet.iter_rows()
            for cell in row
            if cell.coordinate != "A1" and cell.value is not None
        ]
        self.assertEqual(payload, [])


if __name__ == "__main__":
    unittest.main()
