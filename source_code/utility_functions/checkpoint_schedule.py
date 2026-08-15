"""Validated user-defined checkpoint scheduling.

This module owns only the immutable scheduling policy.  Periodic checkpoint
selection and HDF5 persistence remain in ``utility_functions.checkpoint``;
the simulation loop will later arbitrate the two independent trigger sources.
"""

from dataclasses import dataclass
from pathlib import Path
import warnings

import numpy as np
from openpyxl import load_workbook


USER_CHECKPOINTS_INPUT = "USER_CHECKPOINTS"
CHECKPOINTS_SHEET = "CHECKPOINTS"
VALID_SCHEDULED_TRIGGERS = frozenset(("requested", "final"))


@dataclass(frozen=True)
class CheckpointBoundary:
    """One exact simulation time at which a checkpoint must be written."""

    time: float
    trigger: str

    def __post_init__(self):
        if self.trigger not in VALID_SCHEDULED_TRIGGERS:
            raise ValueError(
                f"Invalid scheduled checkpoint trigger {self.trigger!r}."
            )


@dataclass(frozen=True)
class CheckpointSchedule:
    """Immutable normalized checkpoint schedule derived from user input."""

    user_enabled: bool
    boundaries: tuple[CheckpointBoundary, ...]


def parse_user_checkpoints_flag(transient_input):
    """Return the explicit ``USER_CHECKPOINTS`` Boolean value.

    Old input workbooks do not contain this row and therefore default to
    ``False``.  New workbooks accept Excel Booleans and the case-insensitive
    strings ``TRUE`` and ``FALSE`` only.  Numeric aliases and informal values
    such as ``ON`` or ``yes`` are deliberately rejected.
    """

    if USER_CHECKPOINTS_INPUT not in transient_input:
        return False

    raw_value = transient_input[USER_CHECKPOINTS_INPUT]
    if isinstance(raw_value, (bool, np.bool_)):
        return bool(raw_value)

    if isinstance(raw_value, str):
        normalized = raw_value.strip().casefold()
        if normalized == "true":
            return True
        if normalized == "false":
            return False

    raise ValueError(
        f"{USER_CHECKPOINTS_INPUT} must be the Boolean TRUE or FALSE."
    )


def load_checkpoint_schedule(
    workbook_path,
    transient_input,
    *,
    epsilon=1.0e-6,
):
    """Load and normalize the optional ``CHECKPOINTS`` worksheet.

    The worksheet contract is deliberately narrow: cell ``A1`` contains
    ``Time (s)`` and requested times occupy column A from row 2 onward.
    Empty rows are ignored.  Populated cells outside column A are rejected
    when user scheduling is enabled so accidental layout changes cannot be
    interpreted silently.
    """

    workbook_path = Path(workbook_path)
    if not workbook_path.is_file():
        raise FileNotFoundError(
            "The checkpoint schedule workbook does not exist: "
            f"{workbook_path!s}."
        )

    user_enabled = parse_user_checkpoints_flag(transient_input)
    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=False,
    )
    try:
        sheet_present = CHECKPOINTS_SHEET in workbook.sheetnames
        if not sheet_present:
            return build_checkpoint_schedule(
                transient_input,
                (),
                sheet_present=False,
                epsilon=epsilon,
            )

        worksheet = workbook[CHECKPOINTS_SHEET]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    # Everything except A1 is payload.  In disabled mode it is intentionally
    # ignored, but a warning is emitted when any value is present.
    payload_values = [
        value
        for row_index, row in enumerate(rows, start=1)
        for column_index, value in enumerate(row, start=1)
        if (row_index, column_index) != (1, 1) and value is not None
    ]
    if not user_enabled:
        return build_checkpoint_schedule(
            transient_input,
            payload_values,
            sheet_present=True,
            epsilon=epsilon,
        )

    header = rows[0][0] if rows and rows[0] else None
    if not isinstance(header, str) or header.strip() != "Time (s)":
        raise ValueError(
            "CHECKPOINTS!A1 must contain exactly 'Time (s)'."
        )

    unexpected_values = [
        value
        for row in rows
        for value in row[1:]
        if value is not None
    ]
    if unexpected_values:
        raise ValueError(
            "CHECKPOINTS must contain values only in column A."
        )

    checkpoint_times = [
        row[0]
        for row in rows[1:]
        if row and row[0] is not None
    ]
    return build_checkpoint_schedule(
        transient_input,
        checkpoint_times,
        sheet_present=True,
        epsilon=epsilon,
    )


def build_checkpoint_schedule(
    transient_input,
    checkpoint_times=(),
    *,
    sheet_present=True,
    epsilon=1.0e-6,
):
    """Validate and normalize user-defined checkpoint boundaries.

    When ``USER_CHECKPOINTS`` is disabled, the returned schedule is empty;
    periodic checkpoints are intentionally unaffected.  When it is enabled,
    requested times are sorted, values closer than ``epsilon`` are collapsed,
    and ``TEND`` is appended as an automatic ``final`` boundary.  An explicit
    requested value at ``TEND`` retains the higher-priority ``requested``
    trigger instead.
    """

    user_enabled = parse_user_checkpoints_flag(transient_input)
    raw_times = _materialize_times(checkpoint_times)

    if not user_enabled:
        if sheet_present and raw_times:
            warnings.warn(
                "Values in CHECKPOINTS are ignored because "
                "USER_CHECKPOINTS is FALSE.",
                UserWarning,
                stacklevel=2,
            )
        return CheckpointSchedule(user_enabled=False, boundaries=())

    if not sheet_present:
        raise ValueError(
            "The CHECKPOINTS sheet is required when USER_CHECKPOINTS is TRUE."
        )

    epsilon = _positive_finite_float(epsilon, "epsilon")
    tend = _positive_finite_float(
        transient_input.get("TEND"),
        "TEND",
    )
    requested_times = _validate_requested_times(raw_times, tend)

    sorted_times = sorted(requested_times)
    if requested_times != sorted_times:
        warnings.warn(
            "Requested checkpoint times were sorted in increasing order.",
            UserWarning,
            stacklevel=2,
        )

    unique_times = _collapse_near_duplicates(sorted_times, epsilon)
    boundaries = [
        CheckpointBoundary(time=value, trigger="requested")
        for value in unique_times
    ]

    if boundaries and tend - boundaries[-1].time <= epsilon:
        boundaries[-1] = CheckpointBoundary(
            time=tend,
            trigger="requested",
        )
    else:
        boundaries.append(CheckpointBoundary(time=tend, trigger="final"))

    return CheckpointSchedule(
        user_enabled=True,
        boundaries=tuple(boundaries),
    )


def next_checkpoint_boundary(schedule, current_time, epsilon=1.0e-6):
    """Return the first boundary strictly after the restored/current time."""

    if not isinstance(schedule, CheckpointSchedule):
        raise TypeError("schedule must be a CheckpointSchedule instance.")

    current_time = _finite_float(current_time, "current_time")
    epsilon = _positive_finite_float(epsilon, "epsilon")
    threshold = current_time + epsilon
    for boundary in schedule.boundaries:
        if boundary.time > threshold:
            return boundary
    return None


def checkpoint_boundary_due(schedule, current_time, epsilon=1.0e-6):
    """Return the scheduled boundary matching ``current_time``, if any."""

    if not isinstance(schedule, CheckpointSchedule):
        raise TypeError("schedule must be a CheckpointSchedule instance.")

    current_time = _finite_float(current_time, "current_time")
    epsilon = _positive_finite_float(epsilon, "epsilon")
    for boundary in schedule.boundaries:
        if abs(boundary.time - current_time) <= epsilon:
            return boundary
    return None


def _materialize_times(values):
    if values is None:
        return []
    if isinstance(values, (str, bytes)):
        return [values]
    try:
        return list(values)
    except TypeError as exc:
        raise ValueError(
            "CHECKPOINTS values must be an iterable of finite numeric times."
        ) from exc


def _validate_requested_times(values, tend):
    result = []
    for raw_value in values:
        if isinstance(raw_value, (bool, np.bool_)):
            raise ValueError(
                "Every CHECKPOINTS value must be a finite numeric time."
            )
        try:
            value = float(raw_value)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                "Every CHECKPOINTS value must be a finite numeric time."
            ) from exc
        if not np.isfinite(value):
            raise ValueError(
                "Every CHECKPOINTS value must be a finite numeric time."
            )
        if value <= 0.0:
            raise ValueError(
                "Every CHECKPOINTS time must be greater than zero."
            )
        if value > tend:
            raise ValueError("A CHECKPOINTS time cannot exceed TEND.")
        result.append(value)
    return result


def _collapse_near_duplicates(values, epsilon):
    unique = []
    duplicate_found = False
    for value in values:
        if unique and value - unique[-1] <= epsilon:
            duplicate_found = True
            continue
        unique.append(value)

    if duplicate_found:
        warnings.warn(
            "duplicate checkpoint times within the accepted tolerance "
            "were collapsed.",
            UserWarning,
            stacklevel=3,
        )
    return unique


def _positive_finite_float(value, label):
    result = _finite_float(value, label)
    if result <= 0.0:
        raise ValueError(f"{label} must be greater than zero.")
    return result


def _finite_float(value, label):
    if isinstance(value, (bool, np.bool_)):
        raise ValueError(f"{label} must be a finite numeric value.")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be a finite numeric value.") from exc
    if not np.isfinite(result):
        raise ValueError(f"{label} must be a finite numeric value.")
    return result
