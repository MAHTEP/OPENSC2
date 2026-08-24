"""HDF5 checkpoint persistence and recovery utilities.

The reader returns a detached intermediate representation.  Runtime validation
and state application are separate operations so every incompatibility can be
reported before a live OPENSC2 object is mutated.
"""

import copy
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
from urllib.parse import quote

import h5py
import numpy as np
from openpyxl import load_workbook

from utility_functions.checkpoint_schedule import (
    checkpoint_boundary_due,
)
from utility_functions.utils_global_info import IADAPTIME_VALUES


SCHEMA_VERSION = "1.1"
SUPPORTED_SCHEMA_VERSIONS = frozenset((SCHEMA_VERSION,))
CONTINUATION_PROFILE_VERSION = "1.1"
VALID_TRIGGERS = frozenset(("periodic", "requested", "final"))
DEFAULT_CHECKPOINT_EVERY_N_STEPS = 100
CHECKPOINT_INTERVAL_INPUT = "CHECKPOINT_EVERY_N_STEPS"

_CONTINUATION_IMMUTABLE_TRANSIENT_KEYS = ()
_CONTINUATION_RUN_METADATA_KEYS = ("SIMULATION",)
_CONTINUATION_TIME_POLICY_KEYS = (
    "IADAPTIME",
    "TIME_STEP",
    "STPMIN",
    "STPMAX",
    "MLT_INCREASE",
    "MLT_DECREASE",
    "TIMEREF",
    "TAUREF",
    "TEND",
    "CHECKPOINT_EVERY_N_STEPS",
    "USER_CHECKPOINTS",
)

_CONTINUATION_CONDUCTOR_TIME_POLICY_KEYS = ("ELECTRIC_TIME_STEP",)

_CONTINUATION_CONDUCTOR_DRIVER_KEYS = (
    "I0_OP_MODE",
    "I0_OP_TOT",
)

_CONTINUATION_COMPONENT_DRIVER_KEYS = (
    "IOP_MODE",
    "IOP_INTERPOLATION",
    "IBIFUN",
    "BISS",
    "BOSS",
    "BITR",
    "BOTR",
    "B_INTERPOLATION",
    "B_field_units",
    "IALPHAB",
    "ALPHAB_INTERPOLATION",
    "IQFUN",
    "Q_INTERPOLATION",
    "XQBEG",
    "XQEND",
    "Q0",
    "TQBEG",
    "TQEND",
)

_CONTINUATION_DRIVER_FILE_KEYS = frozenset(
    (
        "EXTERNAL_CURRENT",
        "EXTERNAL_BFIELD",
        "EXTERNAL_ALPHAB",
        "EXTERNAL_HEAT",
    )
)

# The operation workbook contains both mutable driver definitions and fixed
# component settings.  Its fixed semantics are captured from the initialized
# runtime mappings, so hashing the whole workbook would incorrectly reject
# legitimate driver changes.
_CONTINUATION_MIXED_FILE_KEYS = frozenset(("OPERATION",))

# Output diagnostics belong to the new branch rather than to the physical
# model restored from a checkpoint.  Keep them out of the immutable profile;
# recovery mode remains protected by the complete input manifest.
_CONTINUATION_OUTPUT_FILE_KEYS = frozenset(("OUTPUT",))

_CURRENT_FUNCTION = (
    "utility_functions.electric_auxiliary_functions."
    "custom_current_function"
)
_HEAT_FUNCTION = "solid_component.SolidComponent.user_heat_function"

_BALANCE_ATTRIBUTES = (
    "enthalpy_balance",
    "enthalpy_inl",
    "enthalpy_out",
    "E_sol_ini",
    "E_str_ini",
    "E_jk_ini",
)

_OUTPUT_STATE_ATTRIBUTES = (
    "i_save",
    "num_step_save",
    "t_save_left",
)

_OUTPUT_BUFFER_ATTRIBUTES = (
    "store_sd_node",
    "store_sd_gauss",
    "time_evol",
    "time_evol_gauss",
    "time_evol_io",
    "time_evol_max_temperature",
)


class CheckpointValidationError(ValueError):
    """Raised when a simulation is not at a safe checkpoint boundary."""


class CheckpointReadError(ValueError):
    """Raised when a checkpoint cannot be safely read or validated."""


@dataclass(frozen=True)
class InputManifestEntry:
    """One input file recorded in a checkpoint manifest."""

    path: str
    sha256: str
    size: int


@dataclass(frozen=True)
class InputManifestChange:
    """One input file whose saved and current contents differ."""

    checkpoint: InputManifestEntry
    current: InputManifestEntry


@dataclass(frozen=True)
class InputManifestComparison:
    """Detached report comparing saved and current simulation inputs."""

    checkpoint_entries: tuple[InputManifestEntry, ...]
    current_entries: tuple[InputManifestEntry, ...]
    missing: tuple[InputManifestEntry, ...]
    added: tuple[InputManifestEntry, ...]
    modified: tuple[InputManifestChange, ...]

    @property
    def is_match(self):
        """Return ``True`` only when both manifests are identical."""

        return not (self.missing or self.added or self.modified)


@dataclass(frozen=True)
class RestartCompatibilityReport:
    """Non-destructive decision about whether a restart may proceed."""

    mode: str
    is_compatible: bool
    manifest_comparison: InputManifestComparison
    blocking_reasons: tuple[str, ...]
    warnings: tuple[str, ...]
    continuation_comparison: "ContinuationProfileComparison | None"


@dataclass(frozen=True)
class RuntimeRestoreValidationReport:
    """Non-destructive validation of a runtime restore destination."""

    is_valid: bool
    blocking_reasons: tuple[str, ...]
    warnings: tuple[str, ...]


@dataclass(frozen=True)
class ConductorCheckpointData:
    """Detached persistent state of one conductor."""

    identifier: str
    th_method: str
    electric_method: str
    clock: dict
    th_history: dict
    electric: dict
    balances: dict
    components: dict
    output_state: dict


@dataclass(frozen=True)
class ContinuationProfile:
    """Normalized input semantics required by continuation checks."""

    immutable: dict
    time_policy: dict
    drivers: dict


@dataclass(frozen=True)
class ContinuationProfileComparison:
    """Deterministic semantic differences between continuation profiles."""

    is_compatible: bool
    immutable_differences: tuple[str, ...]
    time_policy_differences: tuple[str, ...]
    driver_differences: tuple[str, ...]


@dataclass(frozen=True)
class CheckpointData:
    """Validated checkpoint contents, independent of an open HDF5 file."""

    path: Path
    schema_version: str
    created_utc: str
    trigger: str
    git_commit: str
    input_manifest: tuple[InputManifestEntry, ...]
    continuation_profile: ContinuationProfile
    simulation_time: np.ndarray
    num_step: int
    conductors: dict[str, ConductorCheckpointData]


def read_checkpoint(checkpoint_path):
    """Read and validate a checkpoint without mutating a simulation.

    All HDF5 datasets are copied into ordinary Python values or NumPy arrays
    before the file is closed.  The returned object therefore owns no live
    HDF5 handles.
    """

    checkpoint_path = Path(checkpoint_path)
    if not checkpoint_path.is_file():
        raise CheckpointReadError(
            f"Checkpoint file does not exist: {checkpoint_path!s}."
        )

    try:
        with h5py.File(checkpoint_path, "r") as h5file:
            return _read_checkpoint_file(h5file, checkpoint_path.resolve())
    except CheckpointReadError:
        raise
    except (KeyError, OSError, TypeError, ValueError) as exc:
        raise CheckpointReadError(
            f"Cannot read checkpoint {checkpoint_path!s}: {exc}"
        ) from exc


def compare_input_manifest(checkpoint, input_directory):
    """Compare a checkpoint manifest with the current input directory.

    The function is deliberately diagnostic: it reports every difference but
    does not decide whether a recovery or continuation may proceed.  It also
    does not mutate the detached checkpoint representation.
    """

    if not isinstance(checkpoint, CheckpointData):
        raise TypeError("checkpoint must be a CheckpointData instance.")

    checkpoint_entries = tuple(checkpoint.input_manifest)
    current_entries = _build_input_manifest_entries(input_directory)
    checkpoint_by_path = {entry.path: entry for entry in checkpoint_entries}
    current_by_path = {entry.path: entry for entry in current_entries}

    missing = tuple(
        checkpoint_by_path[path]
        for path in sorted(checkpoint_by_path.keys() - current_by_path.keys())
    )
    added = tuple(
        current_by_path[path]
        for path in sorted(current_by_path.keys() - checkpoint_by_path.keys())
    )
    modified = tuple(
        InputManifestChange(
            checkpoint=checkpoint_by_path[path],
            current=current_by_path[path],
        )
        for path in sorted(checkpoint_by_path.keys() & current_by_path.keys())
        if (
            checkpoint_by_path[path].sha256 != current_by_path[path].sha256
            or checkpoint_by_path[path].size != current_by_path[path].size
        )
    )

    return InputManifestComparison(
        checkpoint_entries=checkpoint_entries,
        current_entries=current_entries,
        missing=missing,
        added=added,
        modified=modified,
    )


def evaluate_restart_compatibility(
    checkpoint,
    input_directory,
    mode="recovery",
    runtime_profile=None,
):
    """Evaluate restart compatibility without mutating runtime state.

    ``recovery`` is accepted only when the current input manifest is exactly
    equal to the one persisted in the checkpoint.  ``continuation`` compares
    the profile stored in the checkpoint with the freshly initialized runtime
    profile.  Time-policy and driver changes are reported but permitted;
    immutable differences block the operation.  Unknown modes are programming
    or user-interface errors and are rejected immediately.
    """

    valid_modes = ("recovery", "continuation")
    if mode not in valid_modes:
        raise ValueError(
            f"Invalid restart mode {mode!r}; expected one of {valid_modes!r}."
        )

    comparison = compare_input_manifest(checkpoint, input_directory)
    blocking_reasons = []
    warnings = []
    continuation_comparison = None

    if mode == "continuation":
        if runtime_profile is None:
            blocking_reasons.append(
                "A current continuation profile is required before "
                "continuation can be evaluated."
            )
        else:
            continuation_comparison = compare_continuation_profiles(
                checkpoint.continuation_profile,
                runtime_profile,
            )
            blocking_reasons.extend(
                f"Immutable continuation input differs: {path}."
                for path in continuation_comparison.immutable_differences
            )
            warnings.extend(
                f"Continuation time-policy input differs: {path}."
                for path in continuation_comparison.time_policy_differences
            )
            warnings.extend(
                f"Continuation driver input differs: {path}."
                for path in continuation_comparison.driver_differences
            )
    else:
        blocking_reasons.extend(
            f"Input file is missing: {entry.path}."
            for entry in comparison.missing
        )
        blocking_reasons.extend(
            f"Unexpected input file was added: {entry.path}."
            for entry in comparison.added
        )
        blocking_reasons.extend(
            f"Input file was modified: {change.checkpoint.path}."
            for change in comparison.modified
        )

    return RestartCompatibilityReport(
        mode=mode,
        is_compatible=not blocking_reasons,
        manifest_comparison=comparison,
        blocking_reasons=tuple(blocking_reasons),
        warnings=tuple(warnings),
        continuation_comparison=continuation_comparison,
    )


def validate_runtime_restore_target(checkpoint, simulation, mode="recovery"):
    """Validate a freshly initialized runtime before applying a checkpoint.

    The function checks identities, methods, component kinds, numerical-state
    shapes, event timelines, and output ownership.  It deliberately performs
    no assignments: every incompatibility is accumulated before the immutable
    report is returned.
    """

    if not isinstance(checkpoint, CheckpointData):
        raise TypeError("checkpoint must be a CheckpointData instance.")

    valid_modes = ("recovery", "continuation")
    if mode not in valid_modes:
        raise ValueError(
            f"Invalid restart mode {mode!r}; expected one of {valid_modes!r}."
        )

    reasons = []
    _validate_fresh_runtime_clock(simulation, reasons)
    if mode == "continuation":
        _validate_continuation_time_policy(
            checkpoint,
            simulation,
            reasons,
        )
    runtime_conductors = list(
        getattr(simulation, "list_of_Conductors", ())
    )
    runtime_by_identifier = {}
    for conductor in runtime_conductors:
        identifier = getattr(conductor, "identifier", None)
        if not identifier:
            reasons.append("Runtime conductor is missing an identifier.")
            continue
        if identifier in runtime_by_identifier:
            reasons.append(
                f"Runtime conductor identifier is duplicated: {identifier!r}."
            )
            continue
        runtime_by_identifier[identifier] = conductor

    checkpoint_ids = set(checkpoint.conductors)
    runtime_ids = set(runtime_by_identifier)
    reasons.extend(
        f"Runtime is missing conductor {identifier!r}."
        for identifier in sorted(checkpoint_ids - runtime_ids)
    )
    reasons.extend(
        f"Runtime contains unexpected conductor {identifier!r}."
        for identifier in sorted(runtime_ids - checkpoint_ids)
    )

    _validate_checkpoint_global_clock(checkpoint, reasons)
    for identifier in sorted(checkpoint_ids & runtime_ids):
        _validate_runtime_conductor_target(
            checkpoint.conductors[identifier],
            runtime_by_identifier[identifier],
            reasons,
            mode=mode,
        )

    return RuntimeRestoreValidationReport(
        is_valid=not reasons,
        blocking_reasons=tuple(reasons),
        warnings=(),
    )


def apply_checkpoint_to_runtime(checkpoint, simulation, mode="recovery"):
    """Apply validated checkpoint state to a freshly initialized runtime.

    This function restores state persisted in the current checkpoint schema
    and synchronizes the fluid and solid primary nodal variables from
    ``SYSVAR``.  It does not
    rebuild derived properties, touch output files, or enter the transient
    loop.

    All replacement values are prepared before the first assignment.  Runtime
    container types are preserved where they are operationally significant;
    in particular, time histories remain lists so the solver can append the
    next completed time.
    """

    valid_modes = ("recovery", "continuation")
    if mode not in valid_modes:
        raise ValueError(
            f"Invalid restart mode {mode!r}; expected one of {valid_modes!r}."
        )

    runtime_profile = (
        build_continuation_profile(simulation)
        if mode == "continuation"
        else None
    )
    compatibility = evaluate_restart_compatibility(
        checkpoint,
        getattr(simulation, "basePath", None),
        mode=mode,
        runtime_profile=runtime_profile,
    )
    if not compatibility.is_compatible:
        details = "\n".join(
            f"- {reason}" for reason in compatibility.blocking_reasons
        )
        raise CheckpointValidationError(
            f"Checkpoint inputs are not compatible with {mode}:\n" + details
        )

    report = validate_runtime_restore_target(
        checkpoint, simulation, mode=mode
    )
    if not report.is_valid:
        details = "\n".join(f"- {reason}" for reason in report.blocking_reasons)
        raise CheckpointValidationError(
            "Runtime restore target is not valid:\n" + details
        )

    runtime_by_identifier = {
        conductor.identifier: conductor
        for conductor in simulation.list_of_Conductors
    }

    prepared_simulation_time = _copy_for_runtime(
        checkpoint.simulation_time,
        simulation.simulation_time,
    )
    prepared_conductors = []
    for identifier in sorted(checkpoint.conductors):
        saved = checkpoint.conductors[identifier]
        runtime = runtime_by_identifier[identifier]
        runtime_components = {
            component.identifier: component
            for component, _kind in _component_inventory(
                runtime, f"conductor {identifier!r}"
            )
        }

        clock_attributes = [
            "cond_time",
            "cond_num_step",
            "EQTEIG",
        ]
        if mode == "recovery":
            clock_attributes.extend(("time_step", "i_event"))

        clock = {
            attribute: _copy_for_runtime(
                saved.clock[attribute], getattr(runtime, attribute)
            )
            for attribute in clock_attributes
        }
        if mode == "continuation":
            checkpoint_time = float(
                np.asarray(saved.clock["cond_time"])[-1]
            )
            clock["i_event"] = _continuation_event_index(
                runtime.events_time,
                checkpoint_time,
            )
        for attribute in ("electric_time", "cond_el_num_step"):
            if attribute in saved.clock:
                clock[attribute] = _copy_for_runtime(
                    saved.clock[attribute], getattr(runtime, attribute)
                )

        th_history = _copy_for_runtime(saved.th_history, runtime.dict_Step)
        component_primary_state = _prepare_component_primary_state(
            runtime, th_history
        )
        electric = {
            attribute: _copy_for_runtime(value, getattr(runtime, attribute))
            for attribute, value in saved.electric.items()
        }
        balances = {
            attribute: _copy_for_runtime(value, getattr(runtime, attribute))
            for attribute, value in saved.balances.items()
        }

        component_energy = {}
        for component_identifier, component_state in saved.components.items():
            if component_state["kind"] != "solid":
                continue
            runtime_component = runtime_components[component_identifier]
            component_energy[component_identifier] = {
                key: _copy_for_runtime(
                    value, runtime_component.dict_node_pt[key]
                )
                for key, value in component_state["energy_history"].items()
            }

        if mode == "continuation":
            output_attributes = _prepare_continuation_output_state(
                runtime,
                checkpoint_time,
            )
        else:
            output_attributes = _prepare_recovery_output_state(
                saved.output_state,
                runtime,
            )
        output_buffers = []
        for owner_path, owner in _runtime_output_owners(runtime):
            saved_owner = _mapping_path(
                saved.output_state["buffers"], owner_path
            )
            for attribute, value in saved_owner.items():
                output_buffers.append(
                    (
                        owner,
                        attribute,
                        _copy_for_runtime(value, getattr(owner, attribute)),
                    )
                )

        prepared_conductors.append(
            (
                runtime,
                clock,
                th_history,
                electric,
                balances,
                runtime_components,
                component_primary_state,
                component_energy,
                output_attributes,
                output_buffers,
            )
        )

    simulation.simulation_time = prepared_simulation_time
    simulation.num_step = int(checkpoint.num_step)
    for (
        runtime,
        clock,
        th_history,
        electric,
        balances,
        runtime_components,
        component_primary_state,
        component_energy,
        output_attributes,
        output_buffers,
    ) in prepared_conductors:
        for attribute, value in clock.items():
            setattr(runtime, attribute, value)
        runtime.dict_Step = th_history
        for owner, values in component_primary_state:
            for key, value in values.items():
                owner.dict_node_pt[key] = value
        for attribute, value in electric.items():
            setattr(runtime, attribute, value)
        for attribute, value in balances.items():
            setattr(runtime, attribute, value)
        for component_identifier, energy_history in component_energy.items():
            node_state = runtime_components[component_identifier].dict_node_pt
            for key, value in energy_history.items():
                node_state[key] = value
        for attribute, value in output_attributes.items():
            setattr(runtime, attribute, value)
        for owner, attribute, value in output_buffers:
            setattr(owner, attribute, value)

    # This flag is deliberately the last mutation. Invalid checkpoints and
    # failures while preparing the replacement state must leave a fresh
    # runtime on the normal t=0 path.
    simulation.restored_from_checkpoint = True


def _prepare_component_primary_state(conductor, th_history):
    """Prepare detached primary nodal state reconstructed from ``SYSVAR``."""

    sysvar = np.asarray(th_history["SYSVAR"])
    if sysvar.ndim != 2 or sysvar.shape[1] < 1:
        raise CheckpointValidationError(
            f"Conductor {conductor.identifier!r} SYSVAR must be a two-"
            "dimensional array with at least one history column."
        )

    try:
        ndf = int(conductor.dict_N_equation["NODOFS"])
        equation_index = conductor.equation_index
    except (AttributeError, KeyError, TypeError, ValueError) as error:
        raise CheckpointValidationError(
            f"Conductor {conductor.identifier!r} lacks a valid runtime "
            "equation mapping."
        ) from error

    if ndf <= 0:
        raise CheckpointValidationError(
            f"Conductor {conductor.identifier!r} NODOFS must be positive."
        )

    prepared = []
    for component in conductor.inventory["FluidComponent"].collection:
        try:
            indices = equation_index[component.identifier]
            values = {
                name: sysvar[getattr(indices, name)::ndf, 0].copy()
                for name in ("velocity", "pressure", "temperature")
            }
            owner = component.coolant
        except (AttributeError, KeyError, TypeError) as error:
            raise CheckpointValidationError(
                f"Fluid component {component.identifier!r} lacks a valid "
                "runtime equation mapping or nodal state."
            ) from error
        _validate_prepared_primary_shapes(component.identifier, owner, values)
        prepared.append((owner, values))

    for component in conductor.inventory["SolidComponent"].collection:
        try:
            values = {
                "temperature": sysvar[
                    equation_index[component.identifier]::ndf, 0
                ].copy()
            }
        except (KeyError, TypeError) as error:
            raise CheckpointValidationError(
                f"Solid component {component.identifier!r} lacks a valid "
                "runtime equation mapping or nodal state."
            ) from error
        _validate_prepared_primary_shapes(
            component.identifier, component, values
        )
        prepared.append((component, values))

    return prepared


def _validate_prepared_primary_shapes(identifier, owner, values):
    """Reject invalid SYSVAR slices before any runtime state is mutated."""

    node_state = getattr(owner, "dict_node_pt", None)
    if not isinstance(node_state, Mapping):
        raise CheckpointValidationError(
            f"Component {identifier!r} lacks a nodal state mapping."
        )

    for key, value in values.items():
        if key not in node_state:
            raise CheckpointValidationError(
                f"Component {identifier!r} nodal state is missing {key!r}."
            )
        expected_shape = np.asarray(node_state[key]).shape
        if value.shape != expected_shape:
            raise CheckpointValidationError(
                f"Component {identifier!r} {key!r} reconstructed from "
                f"SYSVAR has shape {value.shape}, expected {expected_shape}."
            )


def _copy_for_runtime(saved, runtime_template):
    """Deep-copy saved data while retaining mutable runtime container types."""

    if isinstance(runtime_template, Mapping):
        # Legacy schema-1.1 HDF5 groups may expose mapping keys in
        # lexicographic order.  Rebuild shared entries in the order of the
        # already initialized runtime mapping, then append saved-only keys
        # that are legitimately created lazily during the transient.
        restored = {
            key: _copy_for_runtime(saved[key], value)
            for key, value in runtime_template.items()
            if key in saved
        }
        restored.update(
            {
                key: copy.deepcopy(value)
                for key, value in saved.items()
                if key not in runtime_template
            }
        )
        return restored
    if isinstance(runtime_template, list):
        values = saved.tolist() if isinstance(saved, np.ndarray) else list(saved)
        return copy.deepcopy(values)
    if isinstance(runtime_template, tuple):
        values = saved.tolist() if isinstance(saved, np.ndarray) else list(saved)
        return tuple(copy.deepcopy(values))
    if isinstance(runtime_template, np.ndarray):
        return np.array(saved, dtype=runtime_template.dtype, copy=True)
    if isinstance(runtime_template, np.generic):
        return np.asarray(saved, dtype=runtime_template.dtype).item()
    return copy.deepcopy(saved)


def _runtime_output_owners(conductor):
    """Yield runtime output owners using decoded mapping path components."""

    yield ("conductor",), conductor
    for component, _kind in _component_inventory(
        conductor, f"conductor {conductor.identifier!r}"
    ):
        base_path = ("components", component.identifier)
        yield base_path + ("object",), component
        if hasattr(component, "coolant"):
            yield base_path + ("coolant",), component.coolant
        if hasattr(component, "channel"):
            yield base_path + ("channel",), component.channel


def _mapping_path(mapping, path):
    value = mapping
    for key in path:
        value = value[key]
    return value


def _validate_fresh_runtime_clock(simulation, reasons):
    simulation_time = np.asarray(
        getattr(simulation, "simulation_time", ()), dtype=float
    )
    if (
        simulation_time.shape != (1,)
        or not np.isclose(simulation_time[0], 0.0)
    ):
        reasons.append(
            "Runtime simulation must be freshly initialized at time zero."
        )
    if getattr(simulation, "num_step", None) != 0:
        reasons.append("Runtime simulation num_step must be zero before restore.")


def _validate_continuation_time_policy(checkpoint, simulation, reasons):
    """Validate the time policy selected for a continuation run."""

    transient_input = getattr(simulation, "transient_input", None)
    if not isinstance(transient_input, Mapping):
        reasons.append("Continuation transient input is missing or invalid.")
        return

    iadaptime_value = transient_input.get("IADAPTIME")
    try:
        iadaptime_array = np.asarray(iadaptime_value)
        iadaptime_numeric = float(iadaptime_array)
        valid_iadaptime = (
            iadaptime_array.ndim == 0
            and not isinstance(iadaptime_value, (bool, np.bool_))
            and np.issubdtype(iadaptime_array.dtype, np.number)
            and np.isfinite(iadaptime_array)
            and iadaptime_numeric.is_integer()
            and int(iadaptime_numeric) in IADAPTIME_VALUES
        )
    except (TypeError, ValueError):
        valid_iadaptime = False

    iadaptime = None
    if not valid_iadaptime:
        reasons.append(
            "Continuation IADAPTIME must be one of "
            f"{IADAPTIME_VALUES}."
        )
    else:
        iadaptime = int(iadaptime_numeric)
        if iadaptime == -1:
            reasons.append(
                "Continuation IADAPTIME=-1 is not implemented."
            )

    numeric_values = {}
    numeric_names = ["TIME_STEP", "STPMIN", "TEND"]
    if iadaptime in (-2, 1, 2):
        numeric_names.append("STPMAX")
    if iadaptime in (1, 2):
        numeric_names.extend(("MLT_INCREASE", "MLT_DECREASE"))
    if iadaptime == -2:
        numeric_names.extend(("TIMEREF", "TAUREF"))

    for name in numeric_names:
        value = transient_input.get(name)
        try:
            array = np.asarray(value)
            valid = (
                array.ndim == 0
                and not isinstance(value, (bool, np.bool_))
                and np.issubdtype(array.dtype, np.number)
                and np.isfinite(array)
            )
        except (TypeError, ValueError):
            valid = False

        if not valid:
            reasons.append(
                f"Continuation {name} must be a finite numeric scalar."
            )
            continue

        numeric_values[name] = float(array)

    positive_names = ["TIME_STEP", "STPMIN"]
    if iadaptime in (-2, 1, 2):
        positive_names.append("STPMAX")
    if iadaptime in (1, 2):
        positive_names.extend(("MLT_INCREASE", "MLT_DECREASE"))
    if iadaptime == -2:
        positive_names.append("TAUREF")

    for name in positive_names:
        if name in numeric_values and numeric_values[name] <= 0.0:
            reasons.append(f"Continuation {name} must be positive.")

    if (
        iadaptime in (-2, 1, 2)
        and "STPMIN" in numeric_values
        and "STPMAX" in numeric_values
    ):
        if numeric_values["STPMIN"] > numeric_values["STPMAX"]:
            reasons.append(
                "Continuation STPMIN must not exceed STPMAX."
            )
        if (
            "TIME_STEP" in numeric_values
            and not (
                numeric_values["STPMIN"]
                <= numeric_values["TIME_STEP"]
                <= numeric_values["STPMAX"]
            )
        ):
            reasons.append(
                "Continuation TIME_STEP must lie between STPMIN and "
                "STPMAX for adaptive time stepping."
            )

    if "TEND" not in numeric_values or "STPMIN" not in numeric_values:
        return

    checkpoint_time = float(np.asarray(checkpoint.simulation_time)[-1])
    minimum_remaining_time = 1.0e-5 * numeric_values["STPMIN"]
    if numeric_values["TEND"] - checkpoint_time <= minimum_remaining_time:
        reasons.append(
            "Continuation TEND must be later than the checkpoint time by "
            "more than the simulation time tolerance."
        )


def _validate_checkpoint_global_clock(checkpoint, reasons):
    simulation_time = np.asarray(checkpoint.simulation_time)
    if checkpoint.num_step != simulation_time.size - 1:
        reasons.append(
            "Checkpoint global num_step does not equal "
            "len(simulation_time) - 1."
        )

    conductor_times = []
    for conductor in checkpoint.conductors.values():
        cond_time = np.asarray(conductor.clock.get("cond_time", ()))
        if cond_time.ndim == 1 and cond_time.size:
            conductor_times.append(float(cond_time[-1]))
        if conductor.clock.get("cond_num_step") != checkpoint.num_step:
            reasons.append(
                f"Conductor {conductor.identifier!r} checkpoint step does not "
                "match the global checkpoint step."
            )

    if conductor_times and not np.isclose(
        float(simulation_time[-1]),
        max(conductor_times),
        rtol=1.0e-12,
        atol=1.0e-12,
    ):
        reasons.append(
            "Checkpoint global time does not equal the latest conductor time."
        )


def _validate_runtime_conductor_target(saved, runtime, reasons, mode="recovery"):
    label = f"Conductor {saved.identifier!r}"
    runtime_cond_time = np.asarray(
        getattr(runtime, "cond_time", ()), dtype=float
    )
    if (
        runtime_cond_time.shape != (1,)
        or not np.isclose(runtime_cond_time[0], 0.0)
    ):
        reasons.append(f"{label} runtime clock must start at time zero.")
    if getattr(runtime, "cond_num_step", None) != 0:
        reasons.append(f"{label} runtime cond_num_step must be zero before restore.")
    if getattr(runtime, "i_save", None) != 0:
        reasons.append(f"{label} runtime i_save must be zero before restore.")

    runtime_inputs = getattr(runtime, "inputs", {})
    runtime_th_method = (
        runtime_inputs.get("METHOD")
        if isinstance(runtime_inputs, Mapping)
        else None
    )
    runtime_electric_method = (
        runtime_inputs.get("ELECTRIC_METHOD")
        if isinstance(runtime_inputs, Mapping)
        else None
    )
    if runtime_th_method != saved.th_method:
        reasons.append(
            f"{label} TH method differs: checkpoint={saved.th_method!r}, "
            f"runtime={runtime_th_method!r}."
        )
    if runtime_electric_method != saved.electric_method:
        reasons.append(
            f"{label} electric method differs: "
            f"checkpoint={saved.electric_method!r}, "
            f"runtime={runtime_electric_method!r}."
        )

    try:
        runtime_inventory = _component_inventory(runtime, label)
    except CheckpointValidationError as exc:
        reasons.append(str(exc))
        runtime_inventory = []
    runtime_components = {
        component.identifier: (component, kind)
        for component, kind in runtime_inventory
    }
    saved_ids = set(saved.components)
    runtime_ids = set(runtime_components)
    reasons.extend(
        f"{label} is missing component {identifier!r}."
        for identifier in sorted(saved_ids - runtime_ids)
    )
    reasons.extend(
        f"{label} contains unexpected component {identifier!r}."
        for identifier in sorted(runtime_ids - saved_ids)
    )
    for identifier in sorted(saved_ids & runtime_ids):
        saved_kind = saved.components[identifier]["kind"]
        runtime_component, runtime_kind = runtime_components[identifier]
        if saved_kind != runtime_kind:
            reasons.append(
                f"{label} component {identifier!r} kind differs: "
                f"checkpoint={saved_kind!r}, runtime={runtime_kind!r}."
            )
        elif saved_kind == "solid":
            runtime_node_state = getattr(runtime_component, "dict_node_pt", {})
            for key, saved_value in saved.components[identifier][
                "energy_history"
            ].items():
                if key not in runtime_node_state:
                    reasons.append(
                        f"{label} solid component {identifier!r} is missing "
                        f"runtime energy history {key!r}."
                    )
                else:
                    _validate_matching_shape(
                        saved_value,
                        runtime_node_state[key],
                        f"{label} solid component {identifier!r} {key}",
                        reasons,
                    )

    _validate_matching_mapping_shapes(
        saved.th_history,
        getattr(runtime, "dict_Step", None),
        f"{label} TH history",
        reasons,
    )
    _validate_matching_shape(
        saved.clock.get("EQTEIG"),
        getattr(runtime, "EQTEIG", None),
        f"{label} EQTEIG",
        reasons,
    )
    if mode == "recovery":
        _validate_event_timeline(saved, runtime, label, reasons)
    else:
        _validate_continuation_event_timeline(runtime, label, reasons)
    _validate_electric_target(saved, runtime, label, reasons)
    for attribute in _BALANCE_ATTRIBUTES:
        if not hasattr(runtime, attribute):
            reasons.append(f"{label} is missing balance {attribute!r}.")
    _validate_output_target(
        saved,
        runtime,
        runtime_components,
        label,
        reasons,
        mode=mode,
    )


def _validate_matching_mapping_shapes(saved, runtime, label, reasons):
    if not isinstance(runtime, Mapping):
        reasons.append(f"{label} is missing from the runtime.")
        return
    saved_keys = set(saved)
    runtime_keys = set(runtime)
    reasons.extend(
        f"{label} is missing key {key!r}."
        for key in sorted(saved_keys - runtime_keys)
    )
    reasons.extend(
        f"{label} contains unexpected key {key!r}."
        for key in sorted(runtime_keys - saved_keys)
    )
    for key in sorted(saved_keys & runtime_keys):
        _validate_matching_shape(
            saved[key], runtime[key], f"{label}[{key!r}]", reasons
        )


def _validate_matching_shape(saved, runtime, label, reasons):
    if runtime is None:
        reasons.append(f"{label} is missing from the runtime.")
        return
    saved_shape = np.asarray(saved).shape
    runtime_shape = np.asarray(runtime).shape
    if saved_shape != runtime_shape:
        reasons.append(
            f"{label} shape differs: checkpoint={saved_shape!r}, "
            f"runtime={runtime_shape!r}."
        )


def _validate_event_timeline(saved, runtime, label, reasons):
    saved_events = np.asarray(saved.clock.get("events_time"))
    runtime_events = getattr(runtime, "events_time", None)
    if runtime_events is None:
        reasons.append(f"{label} event timeline is missing from the runtime.")
        return
    runtime_events = np.asarray(runtime_events)
    if (
        saved_events.shape != runtime_events.shape
        or not np.allclose(
            saved_events,
            runtime_events,
            rtol=1.0e-12,
            atol=1.0e-12,
        )
    ):
        reasons.append(f"{label} event timeline differs from the checkpoint.")
    saved_i_event = saved.clock.get("i_event")
    if not isinstance(saved_i_event, (int, np.integer)) or not (
        0 <= int(saved_i_event) < saved_events.size
    ):
        reasons.append(f"{label} saved i_event is outside the event timeline.")


def _validate_continuation_event_timeline(runtime, label, reasons):
    """Validate the freshly built event timeline used after continuation."""

    runtime_events = getattr(runtime, "events_time", None)
    if runtime_events is None:
        reasons.append(f"{label} event timeline is missing from the runtime.")
        return

    try:
        runtime_events = np.asarray(runtime_events, dtype=float)
    except (TypeError, ValueError):
        reasons.append(f"{label} runtime event timeline is not numeric.")
        return

    if runtime_events.ndim != 1 or runtime_events.size == 0:
        reasons.append(
            f"{label} runtime event timeline must be a non-empty 1-D array."
        )
        return
    if not np.isfinite(runtime_events).all():
        reasons.append(
            f"{label} runtime event timeline contains non-finite values."
        )
    if runtime_events.size > 1 and np.any(np.diff(runtime_events) <= 0.0):
        reasons.append(
            f"{label} runtime event timeline must be strictly increasing."
        )
    if getattr(runtime, "i_event", None) != 0:
        reasons.append(
            f"{label} runtime event cursor must be zero before continuation."
        )


def _continuation_event_index(events, checkpoint_time):
    """Return the first event strictly after the continuation boundary."""

    events = np.asarray(events, dtype=float)
    next_index = int(np.searchsorted(events, checkpoint_time, side="right"))
    return min(next_index, events.size - 1)


def _continuation_future_output_times(space_save, checkpoint_time):
    """Return diagnostic times strictly after a continuation boundary."""

    values = np.asarray(space_save, dtype=float)
    at_boundary = np.isclose(
        values,
        checkpoint_time,
        rtol=1.0e-12,
        atol=1.0e-12,
    )
    return values[(values > checkpoint_time) & ~at_boundary]


def _validate_continuation_output_schedule(
    runtime,
    checkpoint_time,
    label,
    reasons,
):
    """Validate the fresh spatial-output policy selected by a branch."""

    runtime_space_save = getattr(runtime, "Space_save", None)
    try:
        space_save = np.asarray(runtime_space_save, dtype=float)
    except (TypeError, ValueError):
        reasons.append(f"{label} continuation Space_save is not numeric.")
        return

    if space_save.ndim != 1 or space_save.size == 0:
        reasons.append(
            f"{label} continuation Space_save must be a non-empty 1-D array."
        )
        return
    if not np.isfinite(space_save).all():
        reasons.append(
            f"{label} continuation Space_save contains non-finite values."
        )
        return
    if space_save.size > 1 and np.any(np.diff(space_save) <= 0.0):
        reasons.append(
            f"{label} continuation Space_save must be strictly increasing."
        )
        return
    if _continuation_future_output_times(space_save, checkpoint_time).size == 0:
        reasons.append(
            f"{label} continuation Space_save contains no time after the "
            "checkpoint."
        )

    runtime_num_step_save = getattr(runtime, "num_step_save", None)
    if runtime_num_step_save is None:
        reasons.append(f"{label} continuation num_step_save is missing.")
    else:
        runtime_num_step_save = np.asarray(runtime_num_step_save)
        if (
            runtime_num_step_save.ndim != 1
            or not np.issubdtype(runtime_num_step_save.dtype, np.integer)
        ):
            reasons.append(
                f"{label} continuation num_step_save must be a 1-D "
                "integer array."
            )


def _prepare_continuation_output_state(runtime, checkpoint_time):
    """Prepare branch-local spatial-output counters and future times."""

    future_times = _continuation_future_output_times(
        runtime.Space_save,
        checkpoint_time,
    ).copy()
    runtime_num_step_save = np.asarray(runtime.num_step_save)
    return {
        "Space_save": future_times,
        "i_save": 0,
        "i_save_max": int(future_times.size - 1),
        "num_step_save": np.zeros(
            future_times.shape,
            dtype=runtime_num_step_save.dtype,
        ),
    }


def _validate_recovery_output_schedule(saved_output, runtime, label, reasons):
    """Validate a saved spatial-output schedule against fresh inputs.

    A continuation removes diagnostic times at or before its branch point.
    Checkpoints written later therefore contain a ``num_step_save`` array
    matching only a suffix of the schedule rebuilt from the unchanged input
    files. Strict recovery may restore that suffix, but it must still reject
    malformed arrays and schedules that cannot originate from the runtime
    policy.
    """

    saved_num_step_save = saved_output.get("num_step_save")
    runtime_num_step_save = getattr(runtime, "num_step_save", None)
    runtime_space_save = getattr(runtime, "Space_save", None)

    try:
        saved_num_step_save = np.asarray(saved_num_step_save)
    except (TypeError, ValueError):
        reasons.append(f"{label} saved num_step_save is invalid.")
        return
    try:
        runtime_num_step_save = np.asarray(runtime_num_step_save)
    except (TypeError, ValueError):
        reasons.append(f"{label} runtime num_step_save is invalid.")
        return
    try:
        runtime_space_save = np.asarray(runtime_space_save, dtype=float)
    except (TypeError, ValueError):
        reasons.append(f"{label} Space_save is missing or non-numeric.")
        return

    if (
        saved_num_step_save.ndim != 1
        or saved_num_step_save.size == 0
        or not np.issubdtype(saved_num_step_save.dtype, np.integer)
    ):
        reasons.append(
            f"{label} saved num_step_save must be a non-empty 1-D "
            "integer array."
        )
    if (
        runtime_num_step_save.ndim != 1
        or not np.issubdtype(runtime_num_step_save.dtype, np.integer)
    ):
        reasons.append(
            f"{label} runtime num_step_save must be a 1-D integer array."
        )
    if runtime_space_save.ndim != 1 or runtime_space_save.size == 0:
        reasons.append(
            f"{label} runtime Space_save must be a non-empty 1-D array."
        )
    elif not np.isfinite(runtime_space_save).all():
        reasons.append(
            f"{label} runtime Space_save contains non-finite values."
        )
    elif runtime_space_save.size > 1 and np.any(
        np.diff(runtime_space_save) <= 0.0
    ):
        reasons.append(
            f"{label} runtime Space_save must be strictly increasing."
        )

    if (
        runtime_num_step_save.ndim == 1
        and runtime_space_save.ndim == 1
        and runtime_num_step_save.size != runtime_space_save.size
    ):
        reasons.append(
            f"{label} runtime num_step_save and Space_save shapes differ."
        )
    if (
        saved_num_step_save.ndim == 1
        and runtime_space_save.ndim == 1
        and saved_num_step_save.size > runtime_space_save.size
    ):
        reasons.append(
            f"{label} saved num_step_save cannot be aligned with runtime "
            "Space_save."
        )

    saved_i_save = saved_output.get("i_save")
    valid_i_save = isinstance(saved_i_save, (int, np.integer))
    if valid_i_save and saved_num_step_save.ndim == 1:
        valid_i_save = 0 <= int(saved_i_save) < saved_num_step_save.size
    if valid_i_save and runtime_space_save.ndim == 1:
        valid_i_save = int(saved_i_save) < runtime_space_save.size
    if not valid_i_save:
        reasons.append(
            f"{label} saved i_save is outside runtime Space_save."
        )


def _prepare_recovery_output_state(saved_output, runtime):
    """Prepare recovery counters and any continuation-pruned schedule."""

    output_attributes = {
        attribute: _copy_for_runtime(
            value,
            getattr(runtime, attribute, None),
        )
        for attribute, value in saved_output.items()
        if attribute != "buffers"
    }

    saved_size = np.asarray(saved_output["num_step_save"]).size
    runtime_space_save = np.asarray(runtime.Space_save)
    if saved_size < runtime_space_save.size:
        offset = runtime_space_save.size - saved_size
        output_attributes["Space_save"] = runtime_space_save[offset:].copy()
        output_attributes["i_save_max"] = int(saved_size - 1)

    return output_attributes


def _validate_electric_target(saved, runtime, label, reasons):
    runtime_inputs = getattr(runtime, "inputs", {})
    runtime_active = (
        isinstance(runtime_inputs, Mapping)
        and runtime_inputs.get("I0_OP_MODE") is not None
    )
    saved_active = bool(saved.electric)
    if saved_active != runtime_active:
        reasons.append(
            f"{label} electric activation differs: "
            f"checkpoint={saved_active!r}, runtime={runtime_active!r}."
        )
        return
    if not saved_active:
        return
    for key in ("electric_solution", "electric_solution_steady"):
        _validate_matching_shape(
            saved.electric.get(key),
            getattr(runtime, key, None),
            f"{label} {key}",
            reasons,
        )


def _validate_output_target(
    saved,
    runtime,
    runtime_components,
    label,
    reasons,
    mode="recovery",
):
    saved_output = saved.output_state
    if mode == "continuation":
        checkpoint_time = float(np.asarray(saved.clock["cond_time"])[-1])
        _validate_continuation_output_schedule(
            runtime,
            checkpoint_time,
            label,
            reasons,
        )
    else:
        _validate_recovery_output_schedule(
            saved_output,
            runtime,
            label,
            reasons,
        )

    saved_buffers = saved_output.get("buffers", {})
    saved_conductor_buffer = saved_buffers.get("conductor")
    if not isinstance(saved_conductor_buffer, Mapping):
        reasons.append(f"{label} conductor output buffer is missing.")
    else:
        _validate_output_owner(
            saved_conductor_buffer, runtime, f"{label} output owner conductor", reasons
        )

    saved_component_buffers = saved_buffers.get("components", {})
    if not isinstance(saved_component_buffers, Mapping):
        reasons.append(f"{label} component output buffers are missing.")
        return
    for identifier in sorted(set(saved_component_buffers) & set(runtime_components)):
        saved_owners = saved_component_buffers[identifier]
        component = runtime_components[identifier][0]
        runtime_owners = {"object": component}
        if hasattr(component, "coolant"):
            runtime_owners["coolant"] = component.coolant
        if hasattr(component, "channel"):
            runtime_owners["channel"] = component.channel

        if not isinstance(saved_owners, Mapping):
            reasons.append(
                f"{label} component {identifier!r} output owners are invalid."
            )
            continue
        saved_owner_names = set(saved_owners)
        runtime_owner_names = set(runtime_owners)
        reasons.extend(
            f"{label} component {identifier!r} is missing output owner {name!r}."
            for name in sorted(saved_owner_names - runtime_owner_names)
        )
        reasons.extend(
            f"{label} component {identifier!r} has unexpected output owner {name!r}."
            for name in sorted(runtime_owner_names - saved_owner_names)
        )
        for owner_name in sorted(saved_owner_names & runtime_owner_names):
            _validate_output_owner(
                saved_owners[owner_name],
                runtime_owners[owner_name],
                f"{label} component {identifier!r} output owner {owner_name!r}",
                reasons,
            )


def _validate_output_owner(saved, runtime, label, reasons):
    if not isinstance(saved, Mapping):
        reasons.append(f"{label} buffer is invalid.")
        return
    for attribute in sorted(saved):
        if attribute not in _OUTPUT_BUFFER_ATTRIBUTES:
            reasons.append(f"{label} has unsupported buffer {attribute!r}.")
        elif not hasattr(runtime, attribute):
            reasons.append(f"{label} is missing buffer {attribute!r}.")
        else:
            _validate_restore_container_structure(
                saved[attribute],
                getattr(runtime, attribute),
                f"{label} buffer {attribute!r}",
                reasons,
            )


def _validate_restore_container_structure(saved, runtime, label, reasons):
    """Validate restorable mappings without requiring eager runtime keys.

    Output dictionaries are populated lazily during the transient.  A fresh
    runtime may therefore lack keys that legitimately exist in a checkpoint;
    those keys will be created by ``_copy_for_runtime``.  Runtime-only keys
    remain blocking because they indicate that the current runtime expects
    state that the checkpoint cannot supply.
    """

    saved_is_mapping = isinstance(saved, Mapping)
    runtime_is_mapping = isinstance(runtime, Mapping)
    if saved_is_mapping != runtime_is_mapping:
        reasons.append(f"{label} container type differs from the runtime.")
        return
    if not saved_is_mapping:
        return

    saved_keys = set(saved)
    runtime_keys = set(runtime)
    reasons.extend(
        f"{label} has unexpected runtime key {key!r}."
        for key in sorted(runtime_keys - saved_keys)
    )
    for key in sorted(saved_keys & runtime_keys):
        _validate_restore_container_structure(
            saved[key], runtime[key], f"{label}[{key!r}]", reasons
        )


def _read_checkpoint_file(h5file, checkpoint_path):
    _require_members(
        h5file,
        ("metadata", "simulation", "conductors"),
        "checkpoint root",
    )

    metadata = h5file["metadata"]
    if not isinstance(metadata, h5py.Group):
        raise CheckpointReadError("Checkpoint entry 'metadata' must be a group.")
    _require_attributes(
        metadata,
        ("schema_version", "created_utc", "trigger", "git_commit", "complete"),
        "metadata",
    )

    schema_version = _text(metadata.attrs["schema_version"], "schema_version")
    if schema_version not in SUPPORTED_SCHEMA_VERSIONS:
        raise CheckpointReadError(
            f"Unsupported checkpoint schema {schema_version!r}; "
            f"expected one of {sorted(SUPPORTED_SCHEMA_VERSIONS)!r}."
        )
    if not _strict_bool(metadata.attrs["complete"], "metadata.complete"):
        raise CheckpointReadError("Checkpoint is incomplete.")

    trigger = _text(metadata.attrs["trigger"], "trigger")
    if trigger not in VALID_TRIGGERS:
        raise CheckpointReadError(
            f"Invalid checkpoint trigger {trigger!r}."
        )
    created_utc = _text(metadata.attrs["created_utc"], "created_utc")
    git_commit = _text(metadata.attrs["git_commit"], "git_commit")
    try:
        parsed_created_utc = datetime.fromisoformat(created_utc)
    except ValueError as exc:
        raise CheckpointReadError(
            "Checkpoint creation time is not valid ISO-8601 text."
        ) from exc
    if parsed_created_utc.tzinfo is None:
        raise CheckpointReadError(
            "Checkpoint creation time must include a UTC offset."
        )
    if not git_commit:
        raise CheckpointReadError("Checkpoint git commit metadata is empty.")
    manifest = _read_input_manifest(metadata)
    continuation_profile = _read_continuation_profile(metadata)

    simulation = h5file["simulation"]
    if not isinstance(simulation, h5py.Group):
        raise CheckpointReadError("Checkpoint entry 'simulation' must be a group.")
    _require_members(simulation, ("simulation_time", "num_step"), "simulation")
    simulation_time = _read_finite_time_dataset(
        simulation["simulation_time"], "simulation/simulation_time"
    )
    num_step = _read_non_negative_integer(
        simulation["num_step"], "simulation/num_step"
    )

    conductors_group = h5file["conductors"]
    if not isinstance(conductors_group, h5py.Group):
        raise CheckpointReadError("Checkpoint entry 'conductors' must be a group.")
    if len(conductors_group) == 0:
        raise CheckpointReadError("Checkpoint contains no conductors.")

    conductors = {}
    for group_name in conductors_group:
        conductor = _read_conductor(conductors_group[group_name], group_name)
        if conductor.identifier in conductors:
            raise CheckpointReadError(
                f"Duplicate conductor identifier {conductor.identifier!r}."
            )
        conductors[conductor.identifier] = conductor

    return CheckpointData(
        path=checkpoint_path,
        schema_version=schema_version,
        created_utc=created_utc,
        trigger=trigger,
        git_commit=git_commit,
        input_manifest=manifest,
        continuation_profile=continuation_profile,
        simulation_time=simulation_time,
        num_step=num_step,
        conductors=conductors,
    )


def _read_input_manifest(metadata):
    _require_members(metadata, ("input_manifest",), "metadata")
    group = metadata["input_manifest"]
    if not isinstance(group, h5py.Group):
        raise CheckpointReadError("metadata/input_manifest must be a group.")
    _require_members(group, ("paths", "sha256", "sizes"), "input manifest")

    paths = _read_string_vector(group["paths"], "input manifest paths")
    hashes = _read_string_vector(group["sha256"], "input manifest hashes")
    sizes = np.asarray(group["sizes"][()])
    if sizes.ndim != 1 or not np.issubdtype(sizes.dtype, np.integer):
        raise CheckpointReadError("Input manifest sizes must be a 1-D integer array.")
    if not (len(paths) == len(hashes) == len(sizes)):
        raise CheckpointReadError("Input manifest columns have different lengths.")

    entries = []
    seen_paths = set()
    for path, sha256, size in zip(paths, hashes, sizes):
        candidate = Path(path)
        if (
            not path
            or candidate.is_absolute()
            or ".." in candidate.parts
            or candidate.as_posix() != path
        ):
            raise CheckpointReadError(
                f"Invalid relative path in input manifest: {path!r}."
            )
        if path in seen_paths:
            raise CheckpointReadError(
                f"Duplicate path in input manifest: {path!r}."
            )
        if re.fullmatch(r"[0-9a-f]{64}", sha256) is None:
            raise CheckpointReadError(
                f"Invalid SHA-256 in input manifest for {path!r}."
            )
        if int(size) < 0:
            raise CheckpointReadError(
                f"Negative input size in manifest for {path!r}."
            )
        seen_paths.add(path)
        entries.append(InputManifestEntry(path, sha256, int(size)))
    return tuple(entries)


def _read_continuation_profile(metadata):
    _require_members(
        metadata,
        ("continuation_profile",),
        "metadata",
    )
    profile_group = metadata["continuation_profile"]
    if not isinstance(profile_group, h5py.Group):
        raise CheckpointReadError(
            "metadata/continuation_profile must be a group."
        )
    _require_attributes(
        profile_group,
        ("profile_version",),
        "metadata/continuation_profile",
    )
    profile_version = _text(
        profile_group.attrs["profile_version"],
        "continuation profile version",
    )
    if profile_version != CONTINUATION_PROFILE_VERSION:
        raise CheckpointReadError(
            f"Unsupported continuation profile {profile_version!r}; "
            f"expected {CONTINUATION_PROFILE_VERSION!r}."
        )

    section_names = ("immutable", "time_policy", "drivers")
    _require_members(
        profile_group,
        section_names,
        "metadata/continuation_profile",
    )
    sections = {}
    for section_name in section_names:
        section = _read_node(profile_group[section_name])
        if not isinstance(section, Mapping):
            raise CheckpointReadError(
                "Continuation profile section "
                f"{section_name!r} must be a mapping."
            )
        sections[section_name] = section

    return ContinuationProfile(**sections)


def _read_conductor(group, group_name):
    if not isinstance(group, h5py.Group):
        raise CheckpointReadError(
            f"conductors/{group_name} must be a group."
        )
    _require_attributes(group, ("identifier",), f"conductors/{group_name}")
    identifier = _text(group.attrs["identifier"], "conductor identifier")
    if not identifier or _safe_name(identifier) != group_name:
        raise CheckpointReadError(
            f"Conductor group {group_name!r} does not match identifier "
            f"{identifier!r}."
        )

    required_groups = (
        "metadata",
        "clock",
        "th_history",
        "electric",
        "balances",
        "components",
        "output_state",
    )
    _require_members(group, required_groups, f"conductor {identifier!r}")
    for name in required_groups:
        if not isinstance(group[name], h5py.Group):
            raise CheckpointReadError(
                f"Conductor {identifier!r} entry {name!r} must be a group."
            )

    metadata = group["metadata"]
    _require_attributes(
        metadata,
        ("th_method", "electric_method"),
        f"conductor {identifier!r} metadata",
    )
    clock = _read_mapping(group["clock"])
    _validate_read_clock(clock, identifier)

    th_history = _read_mapping(group["th_history"])
    for key in ("SYSVAR", "SYSLOD"):
        if key not in th_history:
            raise CheckpointReadError(
                f"Conductor {identifier!r} TH history is missing {key!r}."
            )

    electric = _read_mapping(group["electric"])
    if electric:
        for key in ("electric_solution", "electric_solution_steady"):
            if key not in electric:
                raise CheckpointReadError(
                    f"Conductor {identifier!r} electric state is missing {key!r}."
                )
        for key in ("cond_el_num_step", "electric_time"):
            if key not in clock:
                raise CheckpointReadError(
                    f"Conductor {identifier!r} clock is missing {key!r} for "
                    "an active electric state."
                )
        electric_time = np.asarray(clock["electric_time"])
        if electric_time.ndim != 0 or not np.isfinite(electric_time):
            raise CheckpointReadError(
                f"Conductor {identifier!r}: electric_time must be a finite scalar."
            )
        if clock["cond_num_step"] > 0 and not np.isclose(
            float(electric_time),
            np.asarray(clock["cond_time"])[-1],
            rtol=1.0e-12,
            atol=1.0e-12,
        ):
            raise CheckpointReadError(
                f"Conductor {identifier!r}: electric_time must coincide with "
                "cond_time[-1]."
            )

    balances = _read_mapping(group["balances"])
    for key in _BALANCE_ATTRIBUTES:
        if key not in balances:
            raise CheckpointReadError(
                f"Conductor {identifier!r} balances are missing {key!r}."
            )

    components = _read_components(group["components"], identifier)

    output_state = _read_mapping(group["output_state"])
    for key in ("i_save", "num_step_save", "buffers"):
        if key not in output_state:
            raise CheckpointReadError(
                f"Conductor {identifier!r} output state is missing {key!r}."
            )
    buffers = output_state["buffers"]
    if not isinstance(buffers, Mapping):
        raise CheckpointReadError(
            f"Conductor {identifier!r} output buffers must be a mapping."
        )
    component_buffers = buffers.get("components")
    if not isinstance(component_buffers, Mapping):
        raise CheckpointReadError(
            f"Conductor {identifier!r} output buffers are missing components."
        )
    if set(component_buffers) != set(components):
        raise CheckpointReadError(
            f"Conductor {identifier!r} component inventory does not match "
            "its component output buffers."
        )

    return ConductorCheckpointData(
        identifier=identifier,
        th_method=_text(metadata.attrs["th_method"], "TH method"),
        electric_method=_text(
            metadata.attrs["electric_method"], "electric method"
        ),
        clock=clock,
        th_history=th_history,
        electric=electric,
        balances=balances,
        components=components,
        output_state=output_state,
    )


def _read_components(group, conductor_identifier):
    if len(group) == 0:
        raise CheckpointReadError(
            f"Conductor {conductor_identifier!r} contains no components."
        )

    components = {}
    for group_name in group:
        component_group = group[group_name]
        if not isinstance(component_group, h5py.Group):
            raise CheckpointReadError(
                f"Component entry {group_name!r} must be a group."
            )
        _require_attributes(
            component_group,
            ("identifier", "kind"),
            f"component {group_name!r}",
        )
        identifier = _text(
            component_group.attrs["identifier"], "component identifier"
        )
        if not identifier or _safe_name(identifier) != group_name:
            raise CheckpointReadError(
                f"Component group {group_name!r} does not match identifier "
                f"{identifier!r}."
            )
        if identifier in components:
            raise CheckpointReadError(
                f"Duplicate component identifier {identifier!r}."
            )

        kind = _text(component_group.attrs["kind"], "component kind")
        if kind not in ("fluid", "solid"):
            raise CheckpointReadError(
                f"Component {identifier!r} has invalid kind {kind!r}."
            )
        state = _read_mapping(component_group)
        energy_history = state.get("energy_history")
        if kind == "solid":
            if not isinstance(energy_history, Mapping):
                raise CheckpointReadError(
                    f"Solid component {identifier!r} is missing energy_history."
                )
            for key in ("EEXT", "EJHT"):
                if key not in energy_history:
                    raise CheckpointReadError(
                        f"Solid component {identifier!r} energy history is "
                        f"missing {key!r}."
                    )
        elif energy_history is not None:
            raise CheckpointReadError(
                f"Fluid component {identifier!r} cannot contain energy_history."
            )

        components[identifier] = {"kind": kind, **state}
    return components


def _validate_read_clock(clock, identifier):
    for key in ("cond_time", "cond_num_step", "time_step", "EQTEIG", "i_event", "events_time"):
        if key not in clock:
            raise CheckpointReadError(
                f"Conductor {identifier!r} clock is missing {key!r}."
            )
    try:
        cond_time = _finite_time_history(
            clock["cond_time"], f"conductor {identifier!r}.cond_time"
        )
    except CheckpointValidationError as exc:
        raise CheckpointReadError(str(exc)) from exc
    cond_num_step = _python_non_negative_integer(
        clock["cond_num_step"],
        f"conductor {identifier!r}.cond_num_step",
    )
    if cond_num_step != len(cond_time) - 1:
        raise CheckpointReadError(
            f"Conductor {identifier!r}: cond_num_step ({cond_num_step}) must equal "
            f"len(cond_time) - 1 ({len(cond_time) - 1})."
        )
    time_step = np.asarray(clock["time_step"])
    if time_step.ndim != 0 or not np.isfinite(time_step) or float(time_step) < 0.0:
        raise CheckpointReadError(
            f"Conductor {identifier!r}: time_step must be a finite "
            "non-negative scalar."
        )


def _read_mapping(group):
    result = {}
    for name in group:
        node = group[name]
        original_name = _text(
            node.attrs.get("original_name", name),
            f"original name for {node.name}",
        )
        if original_name in result:
            raise CheckpointReadError(
                f"Duplicate reconstructed key {original_name!r} in {group.name}."
            )
        result[original_name] = _read_node(node)
    return result


def _read_node(node):
    if isinstance(node, h5py.Dataset):
        value = node[()]
        if isinstance(value, bytes):
            return value.decode("utf-8")
        if isinstance(value, np.ndarray):
            if value.dtype.kind in ("S", "O"):
                decoded = np.empty(value.shape, dtype=object)
                for index in np.ndindex(value.shape):
                    item = value[index]
                    decoded[index] = (
                        item.decode("utf-8") if isinstance(item, bytes) else item
                    )
                return decoded
            return value.copy()
        return value.item() if isinstance(value, np.generic) else value

    if not isinstance(node, h5py.Group):
        raise CheckpointReadError(f"Unsupported HDF5 object at {node.name}.")
    python_type = _text(
        node.attrs.get("python_type", "mapping"),
        f"python_type for {node.name}",
    )
    if python_type == "none":
        if len(node):
            raise CheckpointReadError(f"None node {node.name} must be empty.")
        return None
    values = _read_mapping(node)
    if python_type == "mapping":
        return values
    if python_type in ("list", "tuple"):
        expected_keys = [str(index) for index in range(len(values))]
        if set(values) != set(expected_keys):
            raise CheckpointReadError(
                f"Sequence node {node.name} has invalid indices."
            )
        sequence = [values[key] for key in expected_keys]
        return tuple(sequence) if python_type == "tuple" else sequence
    raise CheckpointReadError(
        f"Unsupported python_type {python_type!r} at {node.name}."
    )


def _require_members(group, names, label):
    missing = [name for name in names if name not in group]
    if missing:
        raise CheckpointReadError(
            f"{label} is missing required entries: {', '.join(missing)}."
        )


def _require_attributes(item, names, label):
    missing = [name for name in names if name not in item.attrs]
    if missing:
        raise CheckpointReadError(
            f"{label} is missing required attributes: {', '.join(missing)}."
        )


def _text(value, label):
    if isinstance(value, bytes):
        value = value.decode("utf-8")
    if not isinstance(value, str):
        raise CheckpointReadError(f"{label} must be text.")
    return value


def _strict_bool(value, label):
    if not isinstance(value, (bool, np.bool_)):
        raise CheckpointReadError(f"{label} must be boolean.")
    return bool(value)


def _read_string_vector(dataset, label):
    if not isinstance(dataset, h5py.Dataset):
        raise CheckpointReadError(f"{label} must be a dataset.")
    values = np.asarray(dataset[()])
    if values.ndim != 1:
        raise CheckpointReadError(f"{label} must be a 1-D array.")
    return tuple(_text(value, label) for value in values)


def _read_finite_time_dataset(dataset, label):
    if not isinstance(dataset, h5py.Dataset):
        raise CheckpointReadError(f"{label} must be a dataset.")
    try:
        return _finite_time_history(dataset[()], label).copy()
    except CheckpointValidationError as exc:
        raise CheckpointReadError(str(exc)) from exc


def _read_non_negative_integer(dataset, label):
    if not isinstance(dataset, h5py.Dataset):
        raise CheckpointReadError(f"{label} must be a dataset.")
    return _python_non_negative_integer(dataset[()], label)


def _python_non_negative_integer(value, label):
    array = np.asarray(value)
    if array.ndim != 0 or not np.issubdtype(array.dtype, np.integer):
        raise CheckpointReadError(f"{label} must be an integer scalar.")
    result = int(array)
    if result < 0:
        raise CheckpointReadError(f"{label} cannot be negative.")
    return result


def checkpoint_interval(transient_input):
    """Return the configured periodic checkpoint interval.

    A missing input keeps the historical default of one checkpoint every 100
    global TH iterations.  Zero disables periodic checkpoints.
    """

    raw_value = transient_input.get(
        CHECKPOINT_INTERVAL_INPUT,
        DEFAULT_CHECKPOINT_EVERY_N_STEPS,
    )
    if isinstance(raw_value, (bool, np.bool_)):
        raise ValueError(
            f"{CHECKPOINT_INTERVAL_INPUT} must be a non-negative integer."
        )

    try:
        numeric_value = float(raw_value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"{CHECKPOINT_INTERVAL_INPUT} must be a non-negative integer."
        ) from exc

    if (
        not np.isfinite(numeric_value)
        or numeric_value < 0.0
        or not numeric_value.is_integer()
    ):
        raise ValueError(
            f"{CHECKPOINT_INTERVAL_INPUT} must be a non-negative integer."
        )
    return int(numeric_value)


def write_periodic_checkpoint_if_due(simulation):
    """Write a periodic checkpoint when the global step reaches its interval.

    Return the written path, or ``None`` when periodic checkpointing is
    disabled or the current global iteration is not a checkpoint boundary.
    """

    interval = checkpoint_interval(simulation.transient_input)
    if interval == 0 or simulation.num_step % interval != 0:
        return None

    try:
        checkpoint_dir = simulation.dict_path["Checkpoint_dir"]
    except KeyError as exc:
        raise CheckpointValidationError(
            "Missing simulation.dict_path['Checkpoint_dir']."
        ) from exc

    return write_checkpoint(simulation, checkpoint_dir, trigger="periodic")


def write_checkpoint_if_due(simulation):
    """Write at most one checkpoint for all active trigger sources.

    A due scheduled boundary takes precedence over a coincident periodic
    trigger.  The schedule has already resolved requested-versus-final
    precedence, so its trigger can be forwarded directly to the writer.
    Runtimes created before user scheduling existed have no
    ``checkpoint_schedule`` attribute and retain periodic-only behavior.
    """

    scheduled_boundary = None
    schedule = getattr(simulation, "checkpoint_schedule", None)
    if schedule is not None:
        scheduled_boundary = checkpoint_boundary_due(
            schedule,
            simulation.simulation_time[-1],
            epsilon=getattr(simulation, "epsilon", 1.0e-6),
        )

    if scheduled_boundary is not None:
        trigger = scheduled_boundary.trigger
    else:
        interval = checkpoint_interval(simulation.transient_input)
        if interval == 0 or simulation.num_step % interval != 0:
            return None
        trigger = "periodic"

    try:
        checkpoint_dir = simulation.dict_path["Checkpoint_dir"]
    except KeyError as exc:
        raise CheckpointValidationError(
            "Missing simulation.dict_path['Checkpoint_dir']."
        ) from exc

    return write_checkpoint(
        simulation,
        checkpoint_dir,
        trigger=trigger,
    )


def write_checkpoint(simulation, checkpoint_dir, trigger):
    """Write an atomic HDF5 checkpoint and return its final path.

    The caller must invoke this function only after every conductor has
    completed its current TH iteration, including output updates and the
    local call to ``get_time_step``.
    """

    validate_checkpoint_state(simulation)
    if trigger not in VALID_TRIGGERS:
        raise ValueError(
            f"Invalid checkpoint trigger {trigger!r}; expected one of "
            f"{sorted(VALID_TRIGGERS)}."
        )

    checkpoint_dir = Path(checkpoint_dir)
    checkpoint_dir.mkdir(parents=True, exist_ok=True)
    filename = f"checkpoint_step_{int(simulation.num_step):06d}.h5"
    final_path = checkpoint_dir / filename
    temporary_path = checkpoint_dir / f"{filename}.tmp"

    manifest = build_input_manifest(simulation, excluded_dir=checkpoint_dir)

    with h5py.File(temporary_path, "w") as h5file:
        _write_metadata(h5file, simulation, trigger, manifest)
        _write_simulation_state(h5file, simulation)
        h5file.flush()

    # Mark completion only after the whole payload has been flushed and the
    # first file handle has been closed.
    with h5py.File(temporary_path, "r+") as h5file:
        h5file["metadata"].attrs["complete"] = True
        h5file.flush()

    os.replace(temporary_path, final_path)
    return final_path


def validate_checkpoint_state(simulation):
    """Validate invariants required by the selected checkpoint boundary."""

    simulation_time = _finite_time_history(
        getattr(simulation, "simulation_time", None), "simulation_time"
    )
    if not hasattr(simulation, "num_step"):
        raise CheckpointValidationError("Missing simulation.num_step.")
    if not isinstance(simulation.num_step, (int, np.integer)):
        raise CheckpointValidationError("simulation.num_step must be an integer.")
    if simulation.num_step < 0:
        raise CheckpointValidationError("simulation.num_step cannot be negative.")

    conductors = list(getattr(simulation, "list_of_Conductors", ()))
    if not conductors:
        raise CheckpointValidationError("The simulation has no conductors.")

    identifiers = [getattr(item, "identifier", None) for item in conductors]
    if any(not identifier for identifier in identifiers):
        raise CheckpointValidationError("Every conductor needs an identifier.")
    if len(set(identifiers)) != len(identifiers):
        raise CheckpointValidationError("Conductor identifiers must be unique.")

    for conductor in conductors:
        _validate_conductor(conductor)

    return simulation_time


def build_input_manifest(simulation, excluded_dir=None):
    """Return deterministic SHA-256 entries for files under ``basePath``."""

    base_path = Path(getattr(simulation, "basePath", ""))
    return [
        {
            "path": entry.path,
            "sha256": entry.sha256,
            "size": entry.size,
        }
        for entry in _build_input_manifest_entries(
            base_path, excluded_dir=excluded_dir
        )
    ]


def build_continuation_profile(simulation):
    """Build the normalized semantic profile stored in new checkpoints."""

    transient_input = getattr(simulation, "transient_input", {})
    if not isinstance(transient_input, Mapping):
        raise CheckpointValidationError(
            "simulation.transient_input must be a mapping when present."
        )

    immutable = _build_immutable_profile(simulation, transient_input)
    time_policy = {
        key: copy.deepcopy(transient_input[key])
        for key in _CONTINUATION_TIME_POLICY_KEYS
        if key in transient_input
    }
    conductor_time_policy = _build_conductor_time_policy(simulation)
    if conductor_time_policy:
        time_policy["conductors"] = conductor_time_policy

    return ContinuationProfile(
        immutable=immutable,
        time_policy=time_policy,
        drivers=_build_driver_profile(simulation),
    )


def _build_immutable_profile(simulation, transient_input):
    """Return fixed input semantics that a continuation cannot change."""

    immutable = {
        key: copy.deepcopy(transient_input[key])
        for key in _CONTINUATION_IMMUTABLE_TRANSIENT_KEYS
        if key in transient_input
    }

    excluded_transient_keys = (
        set(_CONTINUATION_TIME_POLICY_KEYS)
        | set(_CONTINUATION_IMMUTABLE_TRANSIENT_KEYS)
        | set(_CONTINUATION_RUN_METADATA_KEYS)
    )
    fixed_transient = _mapping_without_keys(
        transient_input,
        excluded_transient_keys,
    )
    if fixed_transient:
        immutable["transient_input"] = fixed_transient

    environment_inputs = getattr(
        getattr(simulation, "environment", None),
        "inputs",
        None,
    )
    if isinstance(environment_inputs, Mapping):
        immutable["environment"] = {
            "inputs": _detached_mapping(environment_inputs),
        }

    conductors = {}
    for conductor in _identified_conductors(simulation):
        conductor_profile = _immutable_conductor_profile(
            simulation,
            conductor,
        )
        conductors[conductor.identifier] = conductor_profile
    if conductors:
        immutable["conductors"] = conductors

    return immutable


def _immutable_conductor_profile(simulation, conductor):
    """Return fixed conductor, component, and structural-file semantics."""

    profile = {}
    inputs = getattr(conductor, "inputs", None)
    if isinstance(inputs, Mapping):
        excluded_input_keys = set(
            _CONTINUATION_CONDUCTOR_TIME_POLICY_KEYS
        ) | set(_CONTINUATION_CONDUCTOR_DRIVER_KEYS)
        fixed_inputs = _mapping_without_keys(inputs, excluded_input_keys)
        if fixed_inputs:
            profile["inputs"] = fixed_inputs

    operations = getattr(conductor, "operations", None)
    if isinstance(operations, Mapping) and operations:
        profile["operations"] = _detached_mapping(operations)

    static_files = _build_static_file_profile(simulation, conductor)
    if static_files:
        profile["static_files"] = static_files

    components = {}
    for component, kind in _component_inventory(
        conductor,
        f"conductor {conductor.identifier!r}",
    ):
        identifier = getattr(component, "identifier", None)
        if not identifier:
            raise CheckpointValidationError(
                f"Conductor {conductor.identifier!r} contains a component "
                "without an identifier."
            )
        component_profile = {"kind": kind}
        component_inputs = getattr(component, "inputs", None)
        if isinstance(component_inputs, Mapping) and component_inputs:
            component_profile["inputs"] = _detached_mapping(
                component_inputs
            )
        component_operations = getattr(component, "operations", None)
        if isinstance(component_operations, Mapping) and component_operations:
            if kind == "solid":
                fixed_operations = _mapping_without_keys(
                    component_operations,
                    _CONTINUATION_COMPONENT_DRIVER_KEYS,
                )
            else:
                fixed_operations = _detached_mapping(
                    component_operations
                )
            if fixed_operations:
                component_profile["operations"] = fixed_operations
        components[identifier] = component_profile
    profile["components"] = components

    return profile


def _build_static_file_profile(simulation, conductor):
    """Return content identities for input files containing fixed semantics."""

    file_input = getattr(conductor, "file_input", None)
    if not isinstance(file_input, Mapping):
        return {}

    excluded_keys = (
        _CONTINUATION_DRIVER_FILE_KEYS
        | _CONTINUATION_MIXED_FILE_KEYS
        | _CONTINUATION_OUTPUT_FILE_KEYS
    )
    result = {}
    for key in sorted(file_input, key=str):
        if key in excluded_keys:
            continue
        raw_path = file_input[key]
        if raw_path is None or (
            isinstance(raw_path, str)
            and raw_path.strip().lower() in ("", "none", "nan")
        ):
            continue
        result[key] = _input_file_identity(
            simulation,
            raw_path,
            key,
        )
    return result


def _input_file_identity(simulation, raw_path, label):
    """Return a portable path and deterministic content identity."""

    if not isinstance(raw_path, (str, os.PathLike)):
        raise CheckpointValidationError(
            f"Invalid static input path for {label}."
        )

    base_path = Path(getattr(simulation, "basePath", "")).resolve()
    path = Path(raw_path)
    if not path.is_absolute():
        path = base_path / path
    path = path.resolve()
    if not _is_relative_to(path, base_path):
        raise CheckpointValidationError(
            f"Static input {label} must be inside simulation.basePath."
        )
    if not path.is_file():
        raise CheckpointValidationError(
            f"Static input file does not exist: {path!s}."
        )

    if path.suffix.casefold() == ".xlsx":
        identity_payload = _xlsx_semantic_payload(path, label)
        sha256 = hashlib.sha256(identity_payload).hexdigest()
        size = len(identity_payload)
    else:
        sha256 = _sha256(path)
        size = path.stat().st_size

    return {
        "path": path.relative_to(base_path).as_posix(),
        "sha256": sha256,
        "size": size,
    }


def _xlsx_semantic_payload(path, label):
    """Return canonical workbook semantics, excluding ZIP metadata."""

    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
            keep_links=True,
        )
    except Exception as error:
        raise CheckpointValidationError(
            f"Could not read static XLSX input {label}: {path!s}."
        ) from error

    try:
        worksheets = []
        for worksheet in workbook.worksheets:
            cells = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    cells.append(
                        (
                            cell.coordinate,
                            cell.data_type,
                            _canonical_xlsx_value(cell.value),
                        )
                    )
            worksheets.append((worksheet.title, cells))

        canonical = {
            "format": "opensc2-xlsx-cells-v1",
            "worksheets": worksheets,
        }
        return json.dumps(
            canonical,
            ensure_ascii=False,
            allow_nan=False,
            separators=(",", ":"),
        ).encode("utf-8")
    finally:
        workbook.close()


def _canonical_xlsx_value(value):
    """Return a JSON-safe scalar while preserving its relevant type."""

    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat()}
    if isinstance(value, timedelta):
        return {
            "type": "timedelta",
            "seconds": value.total_seconds(),
        }
    if isinstance(value, bytes):
        return {"type": "bytes", "value": value.hex()}
    if isinstance(value, float) and not np.isfinite(value):
        return {"type": "float", "value": repr(value)}
    if isinstance(value, (str, int, float, bool)):
        return value
    return {
        "type": f"{type(value).__module__}.{type(value).__qualname__}",
        "value": str(value),
    }


def _mapping_without_keys(mapping, excluded_keys):
    """Return a sorted detached mapping without the selected semantic keys."""

    excluded_keys = set(excluded_keys)
    return {
        key: copy.deepcopy(mapping[key])
        for key in sorted(mapping, key=str)
        if key not in excluded_keys
    }


def _detached_mapping(mapping):
    """Return a deterministic deep copy of an input mapping."""

    return {
        key: copy.deepcopy(mapping[key])
        for key in sorted(mapping, key=str)
    }


def _normalized_continuation_immutable(immutable):
    """Ignore branch-local output metadata stored by profile version 1.1.

    Existing schema-1.1 checkpoints contain ``SIMULATION`` and the ``OUTPUT``
    workbook identity in their immutable section.  Removing them only while
    building new runtime profiles would make those checkpoints unusable, so
    both sides are normalized non-destructively before comparison.
    """

    normalized = copy.deepcopy(immutable)
    transient_input = normalized.get("transient_input")
    if isinstance(transient_input, Mapping):
        transient_input.pop("SIMULATION", None)
        if not transient_input:
            normalized.pop("transient_input", None)

    conductors = normalized.get("conductors")
    if isinstance(conductors, Mapping):
        for conductor in conductors.values():
            if not isinstance(conductor, Mapping):
                continue
            static_files = conductor.get("static_files")
            if not isinstance(static_files, Mapping):
                continue
            static_files.pop("OUTPUT", None)
            if not static_files:
                conductor.pop("static_files", None)

    return normalized


def compare_continuation_profiles(checkpoint_profile, runtime_profile):
    """Compare saved and current semantics without mutating either profile.

    Differences in the time policy and physical drivers are deliberately
    permitted for continuation runs, but remain explicitly reported.
    Differences in the immutable section are blocking.
    """

    for name, profile in (
        ("checkpoint_profile", checkpoint_profile),
        ("runtime_profile", runtime_profile),
    ):
        if not isinstance(profile, ContinuationProfile):
            raise TypeError(f"{name} must be a ContinuationProfile instance.")

    immutable_differences = tuple(
        _continuation_difference_paths(
            _normalized_continuation_immutable(
                checkpoint_profile.immutable
            ),
            _normalized_continuation_immutable(runtime_profile.immutable),
            "immutable",
        )
    )
    time_policy_differences = tuple(
        _continuation_difference_paths(
            checkpoint_profile.time_policy,
            runtime_profile.time_policy,
            "time_policy",
        )
    )
    driver_differences = tuple(
        _continuation_difference_paths(
            checkpoint_profile.drivers,
            runtime_profile.drivers,
            "drivers",
        )
    )
    return ContinuationProfileComparison(
        is_compatible=not immutable_differences,
        immutable_differences=immutable_differences,
        time_policy_differences=time_policy_differences,
        driver_differences=driver_differences,
    )


def _build_conductor_time_policy(simulation):
    """Return conductor-local integration settings keyed by identifier."""

    result = {}
    for conductor in _identified_conductors(simulation):
        inputs = getattr(conductor, "inputs", {})
        if not isinstance(inputs, Mapping):
            continue
        parameters = _selected_parameters(
            inputs,
            _CONTINUATION_CONDUCTOR_TIME_POLICY_KEYS,
        )
        if parameters:
            result[conductor.identifier] = parameters
    return result


def _build_driver_profile(simulation):
    """Normalize current, magnetic, and heating driver definitions."""

    result = {}
    for conductor in _identified_conductors(simulation):
        components = _solid_component_operations(conductor)
        if not components:
            continue

        driver_families = {}
        current = _current_driver_profile(simulation, conductor, components)
        if current is not None:
            driver_families["current"] = current

        family_builders = (
            ("magnetic_field", _magnetic_field_profile),
            ("magnetic_field_gradient", _magnetic_gradient_profile),
            ("external_heat", _external_heat_profile),
        )
        for family_name, builder in family_builders:
            family_components = {}
            for component, operations in components:
                profile = builder(
                    simulation,
                    conductor,
                    operations,
                )
                if profile is not None:
                    family_components[component.identifier] = profile
            if family_components:
                driver_families[family_name] = {
                    "components": family_components,
                }

        if driver_families:
            result[conductor.identifier] = driver_families
    return result


def _identified_conductors(simulation):
    """Return conductors in deterministic identifier order."""

    conductors = list(getattr(simulation, "list_of_Conductors", ()))
    identified = [
        conductor
        for conductor in conductors
        if getattr(conductor, "identifier", None)
    ]
    return sorted(identified, key=lambda item: item.identifier)


def _solid_component_operations(conductor):
    """Return solid components having operation mappings."""

    components = []
    for component in _collection(conductor, "SolidComponent"):
        identifier = getattr(component, "identifier", None)
        operations = getattr(component, "operations", None)
        if identifier and isinstance(operations, Mapping):
            components.append((component, operations))
    return sorted(components, key=lambda item: item[0].identifier)


def _current_driver_profile(simulation, conductor, components):
    """Return the conductor-global current source and component settings."""

    inputs = getattr(conductor, "inputs", {})
    if not isinstance(inputs, Mapping) or "I0_OP_MODE" not in inputs:
        return None

    mode = inputs["I0_OP_MODE"]
    source = _driver_source(
        simulation,
        conductor,
        mode,
        auxiliary_key="EXTERNAL_CURRENT",
        function_path=_CURRENT_FUNCTION,
        canonical_modes=(0,),
        disabled_modes=(None,),
        auxiliary_modes=(-1,),
        function_modes=(-2,),
        label="current",
    )
    component_profiles = {}
    for component, operations in components:
        if "IOP_MODE" not in operations:
            continue
        keys = ["IOP_MODE"]
        if mode == -1:
            keys.append("IOP_INTERPOLATION")
        component_profiles[component.identifier] = _selected_parameters(
            operations,
            keys,
        )

    return {
        "source": source,
        "parameters": _selected_parameters(
            inputs,
            ("I0_OP_MODE", "I0_OP_TOT"),
        ),
        "components": component_profiles,
    }


def _magnetic_field_profile(simulation, conductor, operations):
    """Return one component magnetic-field driver definition."""

    if "IBIFUN" not in operations:
        return None
    mode = operations["IBIFUN"]
    source = _driver_source(
        simulation,
        conductor,
        mode,
        auxiliary_key="EXTERNAL_BFIELD",
        canonical_modes=(0, 1),
        auxiliary_predicate=lambda value: value is not None and value < 0,
        label="magnetic field",
    )
    keys = ["IBIFUN"]
    if mode is not None and mode < 0:
        keys.extend(("B_INTERPOLATION", "B_field_units"))
    elif mode == 0:
        keys.extend(("BISS", "BOSS"))
    elif mode == 1:
        keys.extend(("BISS", "BOSS", "BITR", "BOTR"))
    return {
        "source": source,
        "parameters": _selected_parameters(operations, keys),
    }


def _magnetic_gradient_profile(simulation, conductor, operations):
    """Return one component magnetic-field-gradient definition."""

    if "IALPHAB" not in operations:
        return None
    mode = operations["IALPHAB"]
    source = _driver_source(
        simulation,
        conductor,
        mode,
        auxiliary_key="EXTERNAL_ALPHAB",
        disabled_modes=(0, None),
        auxiliary_predicate=lambda value: value is not None and value <= -1,
        label="magnetic field gradient",
    )
    keys = ["IALPHAB"]
    if mode is not None and mode <= -1:
        keys.append("ALPHAB_INTERPOLATION")
    return {
        "source": source,
        "parameters": _selected_parameters(operations, keys),
    }


def _external_heat_profile(simulation, conductor, operations):
    """Return one component external-heating driver definition."""

    if "IQFUN" not in operations:
        return None
    mode = operations["IQFUN"]
    source = _driver_source(
        simulation,
        conductor,
        mode,
        auxiliary_key="EXTERNAL_HEAT",
        function_path=_HEAT_FUNCTION,
        disabled_modes=(0, None),
        auxiliary_modes=(-1,),
        function_modes=(-2,),
        canonical_predicate=lambda value: value is not None and value > 0,
        label="external heat",
    )
    keys = ["IQFUN"]
    if mode is not None and mode > 0:
        keys.extend(("XQBEG", "XQEND", "Q0", "TQBEG", "TQEND"))
    elif mode == -1:
        keys.extend(("Q_INTERPOLATION", "TQBEG", "TQEND"))
    return {
        "source": source,
        "parameters": _selected_parameters(operations, keys),
    }


def _driver_source(
    simulation,
    conductor,
    mode,
    *,
    auxiliary_key,
    label,
    function_path=None,
    canonical_modes=(),
    disabled_modes=(),
    auxiliary_modes=(),
    function_modes=(),
    canonical_predicate=None,
    auxiliary_predicate=None,
):
    """Return a detached descriptor for one selected driver source."""

    if mode in disabled_modes:
        return {"kind": "disabled"}
    if mode in canonical_modes or (
        canonical_predicate is not None and canonical_predicate(mode)
    ):
        return {"kind": "canonical_input"}
    if mode in auxiliary_modes or (
        auxiliary_predicate is not None and auxiliary_predicate(mode)
    ):
        return _auxiliary_file_source(simulation, conductor, auxiliary_key)
    if mode in function_modes:
        return {
            "kind": "python_function",
            "callable": function_path,
        }
    raise CheckpointValidationError(
        f"Unsupported {label} mode in continuation profile: {mode!r}."
    )


def _auxiliary_file_source(simulation, conductor, file_input_key):
    """Describe an auxiliary input by portable path and content identity."""

    file_input = getattr(conductor, "file_input", {})
    if not isinstance(file_input, Mapping) or file_input_key not in file_input:
        raise CheckpointValidationError(
            f"Missing conductor.file_input[{file_input_key!r}]."
        )

    raw_path = file_input[file_input_key]
    if not isinstance(raw_path, (str, os.PathLike)) or not str(raw_path).strip():
        raise CheckpointValidationError(
            f"Invalid auxiliary driver path for {file_input_key}."
        )

    base_path = Path(getattr(simulation, "basePath", "")).resolve()
    path = Path(raw_path)
    if not path.is_absolute():
        path = base_path / path
    path = path.resolve()
    if not _is_relative_to(path, base_path):
        raise CheckpointValidationError(
            f"Auxiliary driver {file_input_key} must be inside "
            "simulation.basePath."
        )
    if not path.is_file():
        raise CheckpointValidationError(
            f"Auxiliary driver file does not exist: {path!s}."
        )

    return {
        "kind": "auxiliary_file",
        "path": path.relative_to(base_path).as_posix(),
        "sha256": _sha256(path),
        "size": path.stat().st_size,
    }


def _selected_parameters(mapping, keys):
    """Copy selected existing values while preserving key order."""

    return {
        key: copy.deepcopy(mapping[key])
        for key in keys
        if key in mapping
    }


def _continuation_difference_paths(checkpoint_value, runtime_value, path):
    """Return sorted leaf paths whose detached semantic values differ."""

    if isinstance(checkpoint_value, Mapping) and isinstance(
        runtime_value,
        Mapping,
    ):
        differences = []
        keys = sorted(
            set(checkpoint_value) | set(runtime_value),
            key=str,
        )
        for key in keys:
            child_path = f"{path}.{key}"
            if key not in checkpoint_value:
                differences.extend(
                    _continuation_leaf_paths(
                        runtime_value[key],
                        child_path,
                    )
                )
                continue
            if key not in runtime_value:
                differences.extend(
                    _continuation_leaf_paths(
                        checkpoint_value[key],
                        child_path,
                    )
                )
                continue
            differences.extend(
                _continuation_difference_paths(
                    checkpoint_value[key],
                    runtime_value[key],
                    child_path,
                )
            )
        return differences

    if isinstance(checkpoint_value, Mapping) or isinstance(
        runtime_value,
        Mapping,
    ):
        return [path]

    sequence_types = (list, tuple)
    if isinstance(checkpoint_value, sequence_types) and isinstance(
        runtime_value,
        sequence_types,
    ):
        if len(checkpoint_value) != len(runtime_value):
            return [path]
        differences = []
        for index, (checkpoint_item, runtime_item) in enumerate(
            zip(checkpoint_value, runtime_value)
        ):
            differences.extend(
                _continuation_difference_paths(
                    checkpoint_item,
                    runtime_item,
                    f"{path}[{index}]",
                )
            )
        return differences

    if isinstance(checkpoint_value, sequence_types) or isinstance(
        runtime_value,
        sequence_types,
    ):
        return [path]

    return [] if _continuation_values_equal(
        checkpoint_value,
        runtime_value,
    ) else [path]


def _continuation_leaf_paths(value, path):
    """Return deterministic leaf paths for an added or missing subtree."""

    if isinstance(value, Mapping):
        if not value:
            return [path]
        leaves = []
        for key in sorted(value, key=str):
            leaves.extend(
                _continuation_leaf_paths(
                    value[key],
                    f"{path}.{key}",
                )
            )
        return leaves

    if isinstance(value, (list, tuple)):
        if not value:
            return [path]
        leaves = []
        for index, item in enumerate(value):
            leaves.extend(
                _continuation_leaf_paths(
                    item,
                    f"{path}[{index}]",
                )
            )
        return leaves

    return [path]


def _continuation_values_equal(checkpoint_value, runtime_value):
    """Return a scalar Boolean for ordinary and NumPy semantic values."""

    if isinstance(checkpoint_value, np.ndarray) or isinstance(
        runtime_value,
        np.ndarray,
    ):
        try:
            return bool(
                np.array_equal(
                    np.asarray(checkpoint_value),
                    np.asarray(runtime_value),
                    equal_nan=True,
                )
            )
        except TypeError:
            return bool(
                np.array_equal(
                    np.asarray(checkpoint_value),
                    np.asarray(runtime_value),
                )
            )

    try:
        equal = checkpoint_value == runtime_value
    except (TypeError, ValueError):
        return False
    if isinstance(equal, np.ndarray):
        return bool(np.all(equal))
    return bool(equal)


def _build_input_manifest_entries(input_directory, excluded_dir=None):
    """Build a deterministic, detached manifest for one input directory."""

    base_path = Path(input_directory)
    if not base_path.is_dir():
        raise CheckpointValidationError(
            f"Simulation input directory does not exist: {base_path!s}."
        )

    base_path = base_path.resolve()
    excluded_dir = Path(excluded_dir).resolve() if excluded_dir else None
    manifest = []
    for path in sorted(base_path.rglob("*")):
        if not path.is_file():
            continue
        resolved_path = path.resolve()
        if excluded_dir is not None and _is_relative_to(resolved_path, excluded_dir):
            continue
        manifest.append(
            InputManifestEntry(
                path=path.relative_to(base_path).as_posix(),
                sha256=_sha256(path),
                size=path.stat().st_size,
            )
        )
    return tuple(manifest)


def _validate_conductor(conductor):
    label = f"conductor {conductor.identifier!r}"
    cond_time = _finite_time_history(
        getattr(conductor, "cond_time", None), f"{label}.cond_time"
    )

    cond_num_step = getattr(conductor, "cond_num_step", None)
    if cond_num_step != len(cond_time) - 1:
        raise CheckpointValidationError(
            f"{label}: cond_num_step ({cond_num_step}) must equal "
            f"len(cond_time) - 1 ({len(cond_time) - 1})."
        )

    time_step = getattr(conductor, "time_step", None)
    if time_step is None or not np.isfinite(time_step) or time_step < 0.0:
        raise CheckpointValidationError(
            f"{label}: time_step must be a finite non-negative value."
        )

    if getattr(conductor, "force_next_tstep_flag", None) is not False:
        raise CheckpointValidationError(
            f"{label}: force_next_tstep_flag must be False after get_time_step()."
        )

    events_time = np.asarray(getattr(conductor, "events_time", ()), dtype=float)
    if not np.isfinite(events_time).all():
        raise CheckpointValidationError(
            f"{label}: events_time contains non-finite values."
        )
    i_event = getattr(conductor, "i_event", None)
    if events_time.size:
        if i_event is None or not 0 <= int(i_event) < events_time.size:
            raise CheckpointValidationError(
                f"{label}: i_event is outside the reconstructed event timeline."
            )
    elif i_event not in (None, 0):
        raise CheckpointValidationError(
            f"{label}: i_event must be 0 or None when there are no events."
        )

    for attribute in ("EQTEIG", "dict_Step"):
        if not hasattr(conductor, attribute):
            raise CheckpointValidationError(f"{label}: missing {attribute}.")
    for key in ("SYSVAR", "SYSLOD"):
        if key not in conductor.dict_Step:
            raise CheckpointValidationError(
                f"{label}: dict_Step is missing required key {key!r}."
            )

    electric_active = getattr(conductor, "inputs", {}).get("I0_OP_MODE") is not None
    if electric_active:
        for attribute in (
            "electric_solution",
            "electric_solution_steady",
            "electric_time",
            "cond_el_num_step",
        ):
            if not hasattr(conductor, attribute):
                raise CheckpointValidationError(
                    f"{label}: active electric model is missing {attribute}."
                )
        if cond_num_step > 0 and not np.isclose(
            conductor.electric_time,
            cond_time[-1],
            rtol=1.0e-12,
            atol=1.0e-12,
        ):
            raise CheckpointValidationError(
                f"{label}: electric_time must coincide with cond_time[-1]."
            )

    for attribute in _BALANCE_ATTRIBUTES:
        if not hasattr(conductor, attribute):
            raise CheckpointValidationError(f"{label}: missing {attribute}.")

    for attribute in ("i_save", "num_step_save"):
        if not hasattr(conductor, attribute):
            raise CheckpointValidationError(f"{label}: missing {attribute}.")

    component_inventory = _component_inventory(conductor, label)
    for component, kind in component_inventory:
        if kind != "solid":
            continue
        component_label = getattr(component, "identifier", repr(component))
        node_state = getattr(component, "dict_node_pt", {})
        for key in ("EEXT", "EJHT"):
            if key not in node_state:
                raise CheckpointValidationError(
                    f"{label}, solid component {component_label!r}: missing {key}."
                )


def _write_metadata(h5file, simulation, trigger, manifest):
    metadata = h5file.create_group("metadata")
    metadata.attrs["schema_version"] = SCHEMA_VERSION
    metadata.attrs["created_utc"] = datetime.now(timezone.utc).isoformat()
    metadata.attrs["trigger"] = trigger
    metadata.attrs["git_commit"] = _git_commit()
    metadata.attrs["complete"] = False

    input_manifest = metadata.create_group("input_manifest")
    string_type = h5py.string_dtype(encoding="utf-8")
    input_manifest.create_dataset(
        "paths",
        data=np.asarray([entry["path"] for entry in manifest], dtype=object),
        dtype=string_type,
    )
    input_manifest.create_dataset(
        "sha256",
        data=np.asarray([entry["sha256"] for entry in manifest], dtype=object),
        dtype=string_type,
    )
    input_manifest.create_dataset(
        "sizes", data=np.asarray([entry["size"] for entry in manifest], dtype=np.int64)
    )


    continuation_profile = build_continuation_profile(simulation)
    profile_group = metadata.create_group(
        "continuation_profile",
        track_order=True,
    )
    profile_group.attrs["profile_version"] = CONTINUATION_PROFILE_VERSION
    _write_value(
        profile_group,
        "immutable",
        continuation_profile.immutable,
    )
    _write_value(
        profile_group,
        "time_policy",
        continuation_profile.time_policy,
    )
    _write_value(
        profile_group,
        "drivers",
        continuation_profile.drivers,
    )


def _write_simulation_state(h5file, simulation):
    simulation_group = h5file.create_group("simulation")
    simulation_group.create_dataset(
        "simulation_time", data=np.asarray(simulation.simulation_time, dtype=float)
    )
    simulation_group.create_dataset("num_step", data=int(simulation.num_step))

    conductors_group = h5file.create_group("conductors")
    for conductor in simulation.list_of_Conductors:
        conductor_group = conductors_group.create_group(
            _safe_name(conductor.identifier)
        )
        conductor_group.attrs["identifier"] = conductor.identifier
        _write_conductor(conductor_group, conductor)


def _write_conductor(group, conductor):
    metadata = group.create_group("metadata")
    metadata.attrs["th_method"] = getattr(conductor, "inputs", {}).get(
        "METHOD", "unknown"
    )
    metadata.attrs["electric_method"] = getattr(conductor, "inputs", {}).get(
        "ELECTRIC_METHOD", "unknown"
    )

    clock = group.create_group("clock")
    for attribute in (
        "cond_time",
        "cond_num_step",
        "time_step",
        "EQTEIG",
        "i_event",
        "events_time",
        "cond_el_num_step",
        "electric_time",
    ):
        if hasattr(conductor, attribute):
            _write_value(clock, attribute, getattr(conductor, attribute))

    th_history = group.create_group("th_history")
    for key, value in conductor.dict_Step.items():
        _write_value(th_history, key, value)

    electric = group.create_group("electric")
    for attribute in ("electric_solution", "electric_solution_steady"):
        if hasattr(conductor, attribute):
            _write_value(electric, attribute, getattr(conductor, attribute))

    balances = group.create_group("balances")
    for attribute in _BALANCE_ATTRIBUTES:
        _write_value(balances, attribute, getattr(conductor, attribute))

    components = group.create_group("components")
    for component, kind in _component_inventory(
        conductor, f"conductor {conductor.identifier!r}"
    ):
        identifier = getattr(component, "identifier", component.__class__.__name__)
        component_group = components.create_group(_safe_name(identifier))
        component_group.attrs["identifier"] = identifier
        component_group.attrs["kind"] = kind
        if kind == "solid":
            energy_history = component_group.create_group("energy_history")
            _write_value(energy_history, "EEXT", component.dict_node_pt["EEXT"])
            _write_value(energy_history, "EJHT", component.dict_node_pt["EJHT"])

    output_state = group.create_group("output_state")
    for attribute in _OUTPUT_STATE_ATTRIBUTES:
        if hasattr(conductor, attribute):
            _write_value(output_state, attribute, getattr(conductor, attribute))
    output_buffers = output_state.create_group("buffers")
    for owner_path, owner in _iter_output_owners(conductor):
        owner_group = _require_path(output_buffers, owner_path)
        for attribute in _OUTPUT_BUFFER_ATTRIBUTES:
            if hasattr(owner, attribute):
                _write_value(owner_group, attribute, getattr(owner, attribute))


def _iter_output_owners(conductor):
    yield "conductor", conductor
    for component in _collection(conductor, "all_component"):
        identifier = getattr(component, "identifier", component.__class__.__name__)
        base_path = f"components/{_safe_name(identifier)}"
        yield f"{base_path}/object", component
        if hasattr(component, "coolant"):
            yield f"{base_path}/coolant", component.coolant
        if hasattr(component, "channel"):
            yield f"{base_path}/channel", component.channel


def _write_value(parent, name, value):
    safe_name = _safe_name(name)
    if isinstance(value, Mapping):
        child = parent.create_group(safe_name, track_order=True)
        child.attrs["python_type"] = "mapping"
        child.attrs["original_name"] = str(name)
        for key, item in value.items():
            _write_value(child, key, item)
        return

    if value is None:
        child = parent.create_group(safe_name)
        child.attrs["python_type"] = "none"
        child.attrs["original_name"] = str(name)
        return

    if isinstance(value, (list, tuple)):
        array = np.asarray(value)
        if array.dtype == object:
            child = parent.create_group(safe_name)
            child.attrs["python_type"] = "tuple" if isinstance(value, tuple) else "list"
            child.attrs["original_name"] = str(name)
            for index, item in enumerate(value):
                _write_value(child, str(index), item)
            return
        value = array

    if isinstance(value, str):
        dataset = parent.create_dataset(
            safe_name, data=value, dtype=h5py.string_dtype(encoding="utf-8")
        )
    else:
        array = np.asarray(value)
        if array.dtype == object:
            raise TypeError(
                f"Cannot serialize object-valued checkpoint entry {name!r}."
            )
        if array.dtype.kind in ("U", "S"):
            dataset = parent.create_dataset(
                safe_name,
                data=array.astype(object),
                dtype=h5py.string_dtype(encoding="utf-8"),
            )
        else:
            dataset = parent.create_dataset(safe_name, data=value)
    dataset.attrs["original_name"] = str(name)


def _collection(conductor, key):
    inventory = getattr(conductor, "inventory", {})
    item = inventory.get(key) if isinstance(inventory, Mapping) else None
    return list(getattr(item, "collection", ()))


def _component_inventory(conductor, label):
    fluid_components = _collection(conductor, "FluidComponent")
    solid_components = _collection(conductor, "SolidComponent")
    all_components = _collection(conductor, "all_component")
    if not all_components:
        raise CheckpointValidationError(
            f"{label}: all_component cannot be empty."
        )

    fluid_ids = [id(component) for component in fluid_components]
    solid_ids = [id(component) for component in solid_components]
    all_ids = [id(component) for component in all_components]
    if len(set(fluid_ids)) != len(fluid_ids):
        raise CheckpointValidationError(
            f"{label}: FluidComponent contains duplicate objects."
        )
    if len(set(solid_ids)) != len(solid_ids):
        raise CheckpointValidationError(
            f"{label}: SolidComponent contains duplicate objects."
        )
    if set(fluid_ids) & set(solid_ids):
        raise CheckpointValidationError(
            f"{label}: a component cannot be both fluid and solid."
        )
    if len(set(all_ids)) != len(all_ids):
        raise CheckpointValidationError(
            f"{label}: all_component contains duplicate objects."
        )
    if set(all_ids) != set(fluid_ids) | set(solid_ids):
        raise CheckpointValidationError(
            f"{label}: all_component must contain exactly the fluid and solid "
            "components."
        )

    identifiers = [getattr(component, "identifier", None) for component in all_components]
    if any(not identifier for identifier in identifiers):
        raise CheckpointValidationError(
            f"{label}: every component needs an identifier."
        )
    if len(set(identifiers)) != len(identifiers):
        raise CheckpointValidationError(
            f"{label}: component identifiers must be unique."
        )

    fluid_id_set = set(fluid_ids)
    return [
        (component, "fluid" if id(component) in fluid_id_set else "solid")
        for component in all_components
    ]


def _require_path(parent, path):
    group = parent
    for part in path.split("/"):
        group = group.require_group(part)
    return group


def _finite_time_history(values, label):
    if values is None:
        raise CheckpointValidationError(f"Missing {label}.")
    array = np.asarray(values, dtype=float)
    if array.ndim != 1 or not array.size:
        raise CheckpointValidationError(f"{label} must be a non-empty 1-D history.")
    if not np.isfinite(array).all():
        raise CheckpointValidationError(f"{label} contains non-finite values.")
    if np.any(np.diff(array) < 0.0):
        raise CheckpointValidationError(f"{label} must be non-decreasing.")
    return array


def _sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_name(value):
    return quote(str(value), safe="_-.")


def _is_relative_to(path, parent):
    try:
        path.relative_to(parent)
    except ValueError:
        return False
    return True


def _git_commit():
    source_root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        ("git", "rev-parse", "HEAD"),
        cwd=source_root,
        capture_output=True,
        check=False,
        text=True,
    )
    return result.stdout.strip() if result.returncode == 0 else "unknown"
